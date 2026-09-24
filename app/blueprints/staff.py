from datetime import datetime, time, timedelta
from xml.sax.saxutils import escape as xml_escape

import qrcode
import json
import math
import os
import random
import re
import shutil

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    g,
    has_request_context,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from werkzeug.security import check_password_hash
from werkzeug.utils import secure_filename

from flask_babel import gettext as _
from flask_babel import lazy_gettext as _l
from flask_login import current_user, login_required
from sqlalchemy import event, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from .. import db
from ..backup import write_backup_zip
from ..decorators import roles_required
from ..inventory import check_and_deduct_stock, restore_stock_for_order
from ..rate_limit import is_blocked, record_failure
from ..models import (
    Category,
    Floor,
    Ingredient,
    MenuItem,
    MenuItemIngredient,
    MenuItemChannelPrice,
    OrderChannel,
    Table,
    Order,
    OrderItem,
    OrderStockDeduction,
    CancelledOrderLog,
    Settings,
    User,
    LoginLog,
    floor_display_name,
    generate_order_pin,
    parse_quantity_fields,
    calculate_ppn,
    ORDER_STATUSES,
    ORDER_TYPES,
    ORDER_TYPE_LABELS,
    ORDER_TYPE_DINE_IN,
    ORDER_TYPE_TAKEAWAY,
    ORDER_TYPE_OJOL,
    CHANNEL_PRICING_PERCENT,
    CHANNEL_PRICING_MANUAL,
    PAYMENT_METHODS,
    ROLES,
    ROLE_LABELS,
    ROLE_OWNER,
    ROLE_KASIR,
    ROLE_DAPUR,
    ROLE_PELAYAN,
)

staff_bp = Blueprint("staff", __name__)

# SVG sengaja TIDAK diizinkan - file SVG bisa berisi <script> yang ikut
# jalan di origin aplikasi ini kalau URL-nya dibuka langsung.
ALLOWED_LOGO_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
ALLOWED_SOUND_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a"}

# Platform delivery contoh buat Mode Demo (lihat _seed_demo_data()) -
# markup persentase mengikuti kisaran komisi mitra reguler yang berlaku
# per September 2026 (bukan merchant preferred/strategis). Logo-nya
# BUKAN logo resmi Gojek/Grab/Shopee (itu merek dagang pihak lain, tidak
# diambil dari internet) - cuma lencana inisial warna generik sebagai
# placeholder sampai Owner upload logo resmi hasil kemitraan sungguhan.
DEMO_CHANNELS = [
    ("GoFood", 20, "gofood_demo.png"),
    ("GrabFood", 30, "grabfood_demo.png"),
    ("ShopeeFood", 20, "shopeefood_demo.png"),
]

# Dipakai _seed_demo_orders() buat pisahkan menu jadi "minuman" vs
# "makanan" waktu bikin kombinasi pesanan contoh - harus sinkron sama
# nama kategori di app/demo_data/fixture.json.
DEMO_DRINK_CATEGORIES = {"Kopi", "Non-Kopi"}
DEMO_FOOD_CATEGORIES = {"Makanan Berat", "Snack", "Roti & Sandwich", "Dessert", "Pasta"}

# Nama data (bukan teks UI) - sengaja BUKAN dibungkus _() gettext, sama
# seperti nama channel di DEMO_CHANNELS & label meja "Meja {n}" di atas.
DEMO_FLOOR_NAME = "Lantai 1"
DEMO_PACKAGING_CUP_NAME = "Gelas Cup Minuman"
DEMO_PACKAGING_BOX_NAME = "Kotak Makanan Styrofoam"

# Harus sinkron dengan NOTIFICATION_SOUNDS di app/static/js/notify.js.
# "custom" = nada upload sendiri (lihat notification_sound_file).
NOTIFICATION_SOUND_KEYS = {
    "custom",
    "church_bell",
    "double_bell_loud",
    "fire_alarm",
    "siren_alarm",
    "door_bell_loud",
    "kitchen_alarm_urgent",
    "metal_gong",
    "alarm_clock",
    "emergency_beacon",
    "warning_horn",
}

# Kunci Mode Demo & Reset ke Mode Pabrik - disimpan sebagai HASH (bukan
# teks polos) supaya PIN aslinya tidak kebaca langsung walau seseorang
# buka source code ini. Sengaja hardcode di kode, BUKAN di database/UI,
# karena kedua fitur ini harus tetap cuma bisa dibuka developer (pembuat
# aplikasi) walau nanti aplikasi diserahkan ke pemilik toko lain yang
# login sebagai Owner - pembatasan berbasis role saja tidak cukup karena
# Owner tokonya sendiri JUSTRU yang harus dicegah buka fitur ini. Kedua
# fitur pakai PIN yang sama, tapi status "terbuka"-nya disimpan TERPISAH
# di session (lihat demo_mode_unlock()/factory_reset_unlock()) - buka
# kunci salah satu tidak otomatis membuka yang lain.
# Status "terbuka" PIN developer cuma berlaku sebentar - supaya tidak
# nempel terus di session tablet yang dipakai bareng.
DEV_UNLOCK_TTL_SECONDS = 10 * 60
DEV_PIN_RATE_LIMIT = 5
DEV_PIN_RATE_WINDOW_SECONDS = 15 * 60

DEV_PIN_HASH = "scrypt:32768:8:1$059iCzyBFMYlJpjk$6f23d3dfc930cbff3d0b7668f025d55838e79089776657a059b9d51a99451172040d211c0343e8a4c26a96fd4f10d4f457d6e9bbcb192949892d8222c0f5736d"


def finite_float(value):
    """Dipakai sebagai type= di request.form.get() - sama seperti float
    biasa tapi menolak "nan"/"inf" (float() Python menerima keduanya),
    yang kalau lolos bikin stok/harga/persentase jadi tidak masuk akal
    atau error 500 di tempat lain."""

    number = float(value)
    if not math.isfinite(number):
        raise ValueError(value)
    return number


def digits_int(value):
    """type= untuk nominal Rupiah yang diketik bebas ("50.000", "Rp 50000")
    - ambil digitnya saja. Kosong/tanpa digit -> ValueError (jadi None)."""

    digits = re.sub(r"\D", "", value or "")
    if not digits or len(digits) > 12:
        raise ValueError(value)
    return int(digits)


def _excel_safe(value):
    """Teks dari input user (nama bahan, nama meja, username, dst) yang
    diawali = + - @ bakal dibaca Excel sebagai RUMUS - diberi prefix '
    supaya selalu tampil sebagai teks biasa (cegah formula injection)."""

    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _delete_file_after_commit(path):
    """Jadwalkan penghapusan file upload lama SETELAH transaksi database
    berhasil di-commit, bukan langsung - kalau request-nya gagal di
    tengah jalan (validasi lain gagal, error, rollback), baris di database
    masih menunjuk ke file lama, jadi file itu tidak boleh sudah terhapus
    duluan (logo/foto jadi rusak)."""

    if not has_request_context():
        if os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass
        return
    g.setdefault("files_to_delete_after_commit", []).append(path)


@event.listens_for(Session, "after_commit")
def _run_pending_file_deletes(_session):
    if not has_request_context():
        return
    paths = g.pop("files_to_delete_after_commit", None) or []
    for path in paths:
        if os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass


@event.listens_for(Session, "after_rollback")
def _drop_pending_file_deletes(_session):
    if has_request_context():
        g.pop("files_to_delete_after_commit", None)


def get_settings():
    settings = Settings.query.order_by(Settings.id.asc()).first()

    if not settings:
        settings = Settings(shop_name=current_app.config["CAFE_NAME"])
        db.session.add(settings)
        db.session.commit()

    return settings


def _upload_folder(subfolder):
    folder = os.path.join(current_app.static_folder, "uploads", subfolder)
    os.makedirs(folder, exist_ok=True)
    return folder


def _branding_upload_folder():
    return _upload_folder("branding")


def _save_logo(
    file_field_name,
    obj,
    model_field,
    subfolder="branding",
    filename_base=None,
    allowed_extensions=None,
    invalid_format_message=None,
):
    """Simpan file yang diupload untuk field tertentu (logo_square,
    logo_wide, qris_image, foto profil user, nada notifikasi custom, dll).
    Return pesan error (string) kalau format tidak valid, atau None kalau
    berhasil/tidak ada file baru dipilih. Default-nya validasi ekstensi
    gambar - kirim allowed_extensions buat validasi jenis file lain
    (misal nada notifikasi .mp3/.wav)."""

    logo_file = request.files.get(file_field_name)

    if not logo_file or not logo_file.filename:
        return None

    extension = os.path.splitext(logo_file.filename)[1].lower()
    valid_extensions = allowed_extensions or ALLOWED_LOGO_EXTENSIONS

    if extension not in valid_extensions:
        return invalid_format_message or _("Format gambar harus PNG, JPG, JPEG, atau WEBP.")

    upload_folder = _upload_folder(subfolder)
    old_filename = getattr(obj, model_field)
    new_filename = f"{filename_base or model_field}{extension}"

    if old_filename and old_filename != new_filename:
        _delete_file_after_commit(os.path.join(upload_folder, old_filename))

    logo_file.save(os.path.join(upload_folder, new_filename))
    setattr(obj, model_field, new_filename)

    return None


def _remove_logo(obj, model_field, subfolder="branding"):
    old_filename = getattr(obj, model_field)

    if not old_filename:
        return

    _delete_file_after_commit(os.path.join(_upload_folder(subfolder), old_filename))
    setattr(obj, model_field, None)


# ============================================================
# DASHBOARD
# ============================================================

@staff_bp.route("/")
@login_required
def dashboard():
    active_orders = (
        Order.query.filter(Order.status != "served")
        .order_by(Order.created_at.asc())
        .all()
    )

    today = datetime.now().date()
    start = datetime.combine(today, time.min)
    end = datetime.combine(today, time.max)

    paid_orders_today = Order.query.filter(
        Order.is_paid.is_(True),
        Order.paid_at >= start,
        Order.paid_at <= end,
    ).all()

    unpaid_orders = Order.query.filter_by(is_paid=False).all()
    occupied_table_count = len({
        order.table_id for order in Order.occupying_table_query().all()
        if order.table_id is not None
    })
    kitchen_pending_count = Order.query.filter(
        Order.status.in_(["pending", "processing"])
    ).count()
    low_stock_count = sum(
        1 for ingredient in Ingredient.query.all() if ingredient.is_low_stock
    )

    stats = {
        "active_count": len(active_orders),
        "sales_today": sum(order.total for order in paid_orders_today),
        "transactions_today": len(paid_orders_today),
        "unpaid_count": len(unpaid_orders),
    }

    # Notifikasi kecil di kartu menu, mengikuti alur pelanggan datang ->
    # duduk (meja) -> dapur proses -> bayar di kasir.
    menu_badges = {
        "staff.table_map": occupied_table_count,
        "staff.kitchen": kitchen_pending_count,
        "staff.cashier": stats["unpaid_count"],
        "staff.admin_inventory": low_stock_count,
    }

    return render_template(
        "staff/dashboard.html",
        orders=active_orders,
        stats=stats,
        menu_badges=menu_badges,
    )


# ============================================================
# DENAH MEJA - status kosong/terisi ala pemilihan kursi bioskop
# ============================================================

@staff_bp.route("/tables/map", methods=["GET", "POST"])
@login_required
def table_map():
    is_owner = current_user.role == ROLE_OWNER

    if request.method == "POST":
        if not is_owner:
            abort(403)

        label = request.form.get("label", "").strip()
        floor = request.form.get("floor", type=int)
        code_suffix = request.form.get("code", "").strip().lower()
        # "postab-" selalu jadi awalan kode meja - tamu isi sisanya saja
        # lewat form (lihat table_map.html), jaga-jaga kalau awalan itu
        # keikutan ketik manual juga, jangan sampai dobel.
        if code_suffix.startswith("postab-"):
            code_suffix = code_suffix[len("postab-"):]
        # Kode meja dipakai langsung jadi nama file QR (_generate_table_qr)
        # - batasi ke karakter yang aman untuk nama file di semua OS.
        code_suffix = re.sub(r"[^a-z0-9-]", "", code_suffix)
        code = f"postab-{code_suffix}" if code_suffix else ""

        if label and floor and code_suffix:
            table = Table(label=label, floor=floor, code=code)
            db.session.add(table)
            try:
                db.session.commit()
            except IntegrityError:
                # Kode meja unique - dua meja tidak boleh berbagi kode yang
                # sama (dipakai buat identifikasi QR & URL pesan tamu).
                db.session.rollback()
                flash(
                    _("Kode meja %(code)s sudah dipakai meja lain. Pakai kode lain.", code=code),
                    "danger",
                )
                return redirect(url_for("staff.table_map"))

            try:
                _generate_table_qr(table)
            except OSError:
                # Gagal simpan file QR - jangan tinggalkan meja tanpa QR,
                # batalkan juga baris mejanya supaya tidak nyangkut di DB.
                db.session.delete(table)
                db.session.commit()
                flash(
                    _("Gagal membuat QR untuk meja %(label)s. Coba lagi.", label=label),
                    "danger",
                )
                return redirect(url_for("staff.table_map"))

            flash(_("Meja %(label)s ditambahkan beserta QR code-nya.", label=label), "success")
        else:
            flash(_("Lengkapi label, lantai, dan kode meja."), "danger")

        return redirect(url_for("staff.table_map"))

    tables = Table.query.order_by(Table.floor, Table.label).all()

    # Meja dianggap "terisi" selama masih ada pesanan yang menempatinya
    # (lihat Order.occupying_table_query - belum lunas, atau sudah
    # lunas tapi belum "served").
    active_order_by_table = {}
    for order in (
        Order.occupying_table_query()
        .order_by(Order.created_at.asc())
        .all()
    ):
        # Pesanan Bawa Pulang/Ojek Online (table_id kosong) tidak
        # menempati meja - jangan ikut dihitung sebagai meja terisi.
        if order.table_id is not None:
            active_order_by_table.setdefault(order.table_id, order)

    floors = {}
    for table in tables:
        floors.setdefault(table.floor, []).append({
            "table": table,
            "order": active_order_by_table.get(table.id),
        })

    available_count = len(tables) - len(active_order_by_table)

    return render_template(
        "staff/table_map.html",
        floors=sorted(floors.items()),
        available_floors=Floor.query.order_by(Floor.number).all(),
        available_count=available_count,
        total_count=len(tables),
        can_order=current_user.role in (ROLE_OWNER, ROLE_KASIR, ROLE_PELAYAN),
        can_use_cashier=current_user.role in (ROLE_OWNER, ROLE_KASIR),
        is_owner=is_owner,
    )


@staff_bp.route("/tables/status")
@login_required
def tables_status():
    """Dipoll berkala oleh JS di halaman Meja, buat tahu meja mana yang
    terisi/kosong saat ini - dipakai untuk auto-refresh silent supaya
    tampilan di HP pelayan/kasir selalu sinkron sama data di server
    tanpa perlu reload manual."""

    active_order_by_table = {}
    for order in Order.occupying_table_query().all():
        if order.table_id is not None:
            active_order_by_table.setdefault(order.table_id, order.id)

    tables = [
        {"id": table.id, "order_id": active_order_by_table.get(table.id)}
        for table in Table.query.all()
    ]
    return {"tables": tables}


@staff_bp.route("/tables/<int:table_id>/delete", methods=["POST"])
@roles_required(ROLE_OWNER)
def delete_table(table_id):
    table = Table.query.get_or_404(table_id)

    if table.orders:
        flash(
            _(
                "Meja %(label)s tidak bisa dihapus karena sudah punya riwayat "
                "pesanan (dipakai di laporan/struk). Biarkan saja kalau sudah "
                "tidak dipakai, atau ganti labelnya.",
                label=table.label,
            ),
            "danger",
        )
        return redirect(url_for("staff.table_map"))

    qr_path = os.path.join(current_app.root_path, "static", "qrcodes", f"{table.code}.png")
    if os.path.exists(qr_path):
        try:
            os.remove(qr_path)
        except OSError:
            pass

    label = table.label
    db.session.delete(table)
    db.session.commit()

    flash(_("Meja %(label)s berhasil dihapus.", label=label), "success")
    return redirect(url_for("staff.table_map"))


@staff_bp.route("/tables/<int:table_id>/qr-print")
@roles_required(ROLE_OWNER)
def table_qr_print(table_id):
    table = Table.query.get_or_404(table_id)
    settings = get_settings()

    return render_template("staff/table_qr_print.html", table=table, settings=settings)


def _pdf_safe(text):
    """Font base-14 reportlab (Helvetica/Courier) pakai encoding mirip
    cp1252 - karakter di luar situ (emoji, tanda kutip pintar hasil
    copy-paste, dst) di teks yang diketik user (nama toko, nama meja,
    WiFi, item menu, dll) bikin UnicodeEncodeError dan gagalkan seluruh
    PDF. Ganti karakter yang tidak didukung jadi '?' daripada crash 500."""
    return str(text).encode("cp1252", errors="replace").decode("cp1252")


def _draw_table_qr_card(c, table, settings, page_w, page_h, margin):
    """Gambar 1 kartu QR meja ke canvas reportlab yang sudah ada (dipakai
    bareng table_qr_pdf & floor_qr_pdf) - TIDAK showPage()/save() sendiri,
    supaya pemanggilnya bebas atur kapan pindah halaman (1 meja per
    halaman untuk PDF sekumpulan lantai)."""
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader

    content_w = page_w - 2 * margin

    def center_text(text, y, size=10, bold=False, color=(0.05, 0.3, 0.46)):
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.setFillColorRGB(*color)
        c.drawCentredString(page_w / 2, y, _pdf_safe(text))

    y = page_h - margin

    logo_filename = settings.app_logo_filename
    logo_path = None
    if logo_filename:
        logo_path = os.path.join(current_app.static_folder, "uploads", "branding", logo_filename)
        if not os.path.exists(logo_path):
            logo_path = None

    if logo_path:
        logo_reader = ImageReader(logo_path)
        iw, ih = logo_reader.getSize()
        logo_h = min(14 * mm, content_w * 0.5 * ih / iw)
        logo_w = logo_h * (iw / ih)
        c.drawImage(
            logo_reader, (page_w - logo_w) / 2, y - logo_h,
            width=logo_w, height=logo_h, mask="auto",
        )
        y -= logo_h + 3 * mm
    else:
        center_text(settings.shop_name, y, size=13, bold=True)
        y -= 7 * mm

    c.setDash([1.5, 1.5])
    c.setLineWidth(0.6)
    c.setStrokeColorRGB(0.8, 0.8, 0.8)
    c.line(margin, y, page_w - margin, y)
    c.setDash([])
    y -= 8 * mm

    center_text(table.label, y, size=22, bold=True)
    y -= 6 * mm
    center_text(table.floor_label, y, size=9, color=(0.53, 0.53, 0.53))
    y -= 8 * mm

    # Blok "sambungkan WiFi dulu" - lihat catatan di wifi_name/
    # wifi_password (models.py): aplikasi ini jalan di jaringan LOKAL,
    # jadi QR-nya baru kebuka kalau tamu sudah nyambung ke WiFi toko.
    if settings.wifi_name:
        wifi_h = 11 * mm
        c.setFillColorRGB(0.988, 0.890, 0.796)
        c.roundRect(margin, y - wifi_h, content_w, wifi_h, 2 * mm, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 6.5)
        c.setFillColorRGB(0.71, 0.35, 0.11)
        c.drawCentredString(page_w / 2, y - 4.2 * mm, _("SAMBUNGKAN KE WIFI DULU"))
        wifi_value = settings.wifi_name
        if settings.wifi_password:
            wifi_value += f"  ·  {_('Pass')}: {settings.wifi_password}"
        c.setFont("Helvetica-Bold", 9)
        c.setFillColorRGB(0.05, 0.3, 0.46)
        c.drawCentredString(page_w / 2, y - 8.5 * mm, _pdf_safe(wifi_value))
        y -= wifi_h + 4 * mm

    qr_path = os.path.join(current_app.static_folder, "qrcodes", f"{table.code}.png")
    qr_size = 55 * mm
    qr_x = (page_w - qr_size) / 2
    if os.path.exists(qr_path):
        qr_reader = ImageReader(qr_path)
        c.setStrokeColorRGB(0.9, 0.9, 0.9)
        c.setLineWidth(0.6)
        c.rect(qr_x - 2 * mm, y - qr_size - 2 * mm, qr_size + 4 * mm, qr_size + 4 * mm)
        c.drawImage(qr_reader, qr_x, y - qr_size, width=qr_size, height=qr_size, mask="auto")
    y -= qr_size + 8 * mm

    center_text(_("SCAN UNTUK PESAN"), y, size=9.5, bold=True, color=(0.94, 0.47, 0.16))
    y -= 5 * mm
    center_text(_("Lihat menu & pesan langsung dari HP Anda"), y, size=7.5, color=(0.53, 0.53, 0.53))

    center_text("Orulabs © 2026", margin, size=7, color=(0.6, 0.6, 0.6))


@staff_bp.route("/tables/<int:table_id>/qr-print.pdf")
@roles_required(ROLE_OWNER)
def table_qr_pdf(table_id):
    """Versi PDF dari kartu QR meja (tombol "Cetak ke PDF") - kartu kecil
    ukuran tetap (bukan mengikuti panjang isi kayak struk), jadi cukup
    ditata langsung tanpa hitung tinggi dulu."""
    import io

    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas

    table = Table.query.get_or_404(table_id)
    settings = get_settings()

    page_w = 90 * mm
    page_h = 130 * mm + (12 * mm if settings.wifi_name else 0)
    margin = 8 * mm

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=(page_w, page_h))
    _draw_table_qr_card(c, table, settings, page_w, page_h, margin)
    c.showPage()
    c.save()
    buffer.seek(0)

    filename = f"qr-{table.code}.pdf"
    return send_file(
        buffer, mimetype="application/pdf",
        as_attachment=True, download_name=filename,
    )


@staff_bp.route("/floors/<int:floor_number>/qr-print")
@roles_required(ROLE_OWNER)
def floor_qr_print(floor_number):
    tables = Table.query.filter_by(floor=floor_number).order_by(Table.label).all()
    if not tables:
        abort(404)
    settings = get_settings()

    return render_template(
        "staff/floor_qr_print.html",
        tables=tables, floor_number=floor_number, settings=settings,
    )


@staff_bp.route("/floors/<int:floor_number>/qr-print.pdf")
@roles_required(ROLE_OWNER)
def floor_qr_pdf(floor_number):
    """Versi PDF dari SEMUA kartu QR meja di 1 lantai (tombol "Cetak
    Semua QR" di kartu lantai) - 1 meja per halaman, kartu yang sama
    persis dengan table_qr_pdf lewat _draw_table_qr_card()."""
    import io

    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas

    tables = Table.query.filter_by(floor=floor_number).order_by(Table.label).all()
    if not tables:
        abort(404)
    settings = get_settings()

    page_w = 90 * mm
    page_h = 130 * mm + (12 * mm if settings.wifi_name else 0)
    margin = 8 * mm

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=(page_w, page_h))
    for table in tables:
        _draw_table_qr_card(c, table, settings, page_w, page_h, margin)
        c.showPage()
    c.save()
    buffer.seek(0)

    filename = f"qr-{floor_display_name(floor_number).lower().replace(' ', '-')}.pdf"
    return send_file(
        buffer, mimetype="application/pdf",
        as_attachment=True, download_name=filename,
    )


# ============================================================
# INPUT PESANAN MANUAL (oleh pelayan, untuk tamu tanpa HP)
# ============================================================

@staff_bp.route("/orders/new", methods=["GET", "POST"])
@roles_required(ROLE_OWNER, ROLE_KASIR, ROLE_PELAYAN)
def new_order():
    # Meja yang masih ditempati pesanan lain (lihat Order.occupying_table_query
    # - belum lunas, atau lunas tapi belum "served") dianggap "terisi" -
    # tidak ditawarkan di sini supaya pelayan tidak numpuk pesanan baru
    # di meja yang sudah dipakai tamu lain. Cuma pesanan Makan di Tempat
    # yang punya table_id, jadi ini otomatis tidak kena pengaruh oleh
    # pesanan Bawa Pulang/Ojek Online.
    occupied_table_ids = {
        order.table_id for order in Order.occupying_table_query().all()
        if order.table_id is not None
    }
    tables = (
        Table.query.filter(~Table.id.in_(occupied_table_ids))
        .order_by(Table.floor, Table.label)
        .all()
    )
    categories = Category.query.order_by(Category.order, Category.name).all()
    channels = OrderChannel.query.filter_by(is_active=True).order_by(OrderChannel.sort_order).all()
    preselected_table_id = request.args.get("table_id", type=int)

    if request.method == "POST":
        order_type = request.form.get("order_type", ORDER_TYPE_DINE_IN)
        if order_type not in ORDER_TYPES:
            order_type = ORDER_TYPE_DINE_IN

        table = None
        channel = None

        if order_type == ORDER_TYPE_DINE_IN:
            table_id = request.form.get("table_id", type=int)
            table = Table.query.get(table_id)

            if not table:
                flash(_("Meja tidak valid."), "danger")
                return redirect(url_for("staff.new_order"))

            if table.id in occupied_table_ids:
                flash(_("Meja %(label)s sudah terisi pesanan lain.", label=table.label), "danger")
                return redirect(url_for("staff.new_order"))

        elif order_type == ORDER_TYPE_OJOL:
            channel_id = request.form.get("channel_id", type=int)
            channel = OrderChannel.query.filter_by(id=channel_id, is_active=True).first()

            if not channel:
                flash(_("Platform ojol tidak valid."), "danger")
                return redirect(url_for("staff.new_order"))

        order = Order(
            table_id=table.id if table else None,
            order_type=order_type,
            channel_id=channel.id if channel else None,
            source="staff",
            status="pending",
            pin=generate_order_pin(),
        )
        added_any = False

        for menu_item_id, quantity in parse_quantity_fields(request.form):
            menu_item = MenuItem.query.get(menu_item_id)

            if not menu_item or not menu_item.is_orderable:
                continue

            price = channel.price_for(menu_item) if channel else menu_item.price

            order.items.append(
                OrderItem(
                    menu_item_id=menu_item.id,
                    name_snapshot=menu_item.name,
                    price_snapshot=price,
                    quantity=quantity,
                )
            )
            added_any = True

        if not added_any:
            flash(_("Pilih minimal 1 item."), "warning")
            return redirect(url_for("staff.new_order"))

        stock_error = check_and_deduct_stock(order)
        if stock_error:
            flash(stock_error, "danger")
            return redirect(url_for("staff.new_order"))

        db.session.add(order)
        try:
            db.session.commit()
        except IntegrityError:
            # Jaring pengaman terakhir (dijamin database) - meja ini
            # keburu dapat pesanan lain persis di detik yang sama,
            # lolos dari pengecekan occupied_table_ids di atas.
            db.session.rollback()
            if table:
                flash(_("Meja baru saja terisi pesanan lain. Coba pilih meja lain."), "danger")
            else:
                flash(_("Pesanan gagal disimpan. Coba lagi."), "danger")
            return redirect(url_for("staff.new_order"))

        flash(_("Pesanan untuk %(label)s berhasil dibuat.", label=order.display_label), "success")
        if order_type in (ORDER_TYPE_OJOL, ORDER_TYPE_TAKEAWAY):
            # Tamu/driver biasanya langsung menunggu di kasir - arahkan
            # langsung ke situ supaya kasir bisa langsung proses bayar &
            # cetak, tidak perlu muter dulu lewat Dashboard.
            return redirect(url_for("staff.cashier"))
        return redirect(url_for("staff.dashboard"))

    return render_template(
        "staff/order_new.html",
        tables=tables,
        categories=categories,
        channels=channels,
        preselected_table_id=preselected_table_id,
        ORDER_TYPE_DINE_IN=ORDER_TYPE_DINE_IN,
        ORDER_TYPE_TAKEAWAY=ORDER_TYPE_TAKEAWAY,
        ORDER_TYPE_OJOL=ORDER_TYPE_OJOL,
    )


# ============================================================
# DAPUR - update status pesanan
# ============================================================

@staff_bp.route("/kitchen")
@roles_required(ROLE_OWNER, ROLE_DAPUR)
def kitchen():
    orders = (
        Order.query.filter(Order.status != "served")
        .order_by(Order.created_at.asc())
        .all()
    )

    return render_template("staff/kitchen.html", orders=orders)


@staff_bp.route("/kitchen/status")
@roles_required(ROLE_OWNER, ROLE_DAPUR)
def kitchen_status():
    """Dipoll berkala oleh JS di halaman Dapur buat deteksi "ada yang
    berubah" supaya bisa bunyikan notifikasi suara tanpa reload manual.
    Bukan cuma ID pesanan baru - "tanda tangan" tiap pesanan ikut sertakan
    status & jumlah baris item, supaya pesanan LAMA yang tiba-tiba dapat
    tambahan menu (lihat public.add_to_order - status-nya di-reset ke
    "pending" lagi) juga ke-detect sebagai perubahan, bukan cuma diam-diam
    nyempil tanpa dapur pernah notice."""

    signatures = [
        f"{order.id}:{order.status}:{len(order.items)}"
        for order in Order.query.filter(Order.status != "served").all()
    ]
    return {"order_ids": signatures}


@staff_bp.route("/kitchen/list")
@roles_required(ROLE_OWNER, ROLE_DAPUR)
def kitchen_list():
    """Dipanggil JS lewat AJAX buat refresh daftar pesanan tanpa reload
    halaman penuh - full reload bikin browser reset izin autoplay audio
    jadi notifikasi suara berikutnya kena blokir diam-diam."""

    orders = (
        Order.query.filter(Order.status != "served")
        .order_by(Order.created_at.asc())
        .all()
    )
    return render_template("staff/_kitchen_orders.html", orders=orders)


@staff_bp.route("/pelayan/ready-status")
@roles_required(ROLE_OWNER, ROLE_PELAYAN)
def pelayan_ready_status():
    """Dipoll global (base.html) khusus buat notifikasi suara Pelayan -
    bunyi begitu ada pesanan yang statusnya berubah jadi "Siap Diantar",
    tandanya pelayan harus ambil pesanan itu dari dapur ke meja tamu.
    Daftarnya sendiri ada di halaman Pelayan (lihat waiter()), tempat
    pelayan menandai pesanan "Sudah Diantar"."""

    signatures = [
        f"{order.id}:{order.status}"
        for order in Order.query.filter_by(status="ready").all()
    ]
    return {"order_ids": signatures}


@staff_bp.route("/pelayan")
@roles_required(ROLE_OWNER, ROLE_PELAYAN)
def waiter():
    """Daftar pesanan "Siap Diantar" untuk Pelayan - pelayan ambil dari
    dapur, antar ke meja/tamu, lalu tandai "Sudah Diantar" di sini (baru
    setelah itu meja dianggap kosong lagi)."""

    return render_template("staff/waiter.html", orders=_ready_orders())


@staff_bp.route("/pelayan/list")
@roles_required(ROLE_OWNER, ROLE_PELAYAN)
def waiter_list():
    return render_template("staff/_waiter_orders.html", orders=_ready_orders())


def _ready_orders():
    return (
        Order.query.filter_by(status="ready")
        .order_by(Order.created_at.asc())
        .all()
    )


@staff_bp.route("/orders/<int:order_id>/advance", methods=["POST"])
@roles_required(ROLE_OWNER, ROLE_DAPUR, ROLE_PELAYAN)
def advance_status(order_id):
    order = Order.query.get_or_404(order_id)
    # Pelayan cuma boleh langkah terakhir (Siap -> Sudah Diantar), tahap
    # masak tetap urusan Dapur.
    back_url = url_for("staff.waiter") if current_user.role == ROLE_PELAYAN else url_for("staff.kitchen")
    if current_user.role == ROLE_PELAYAN and request.form.get("from_status") != "ready":
        abort(403)

    # Status yang TERLIHAT di layar waktu tombol ditekan (dikirim form) -
    # dipakai sebagai syarat UPDATE supaya klik dobel / layar Dapur lain
    # yang belum ter-refresh tidak bisa melompati tahap (mis. pending
    # langsung ke ready/served dan meja keburu dianggap kosong).
    from_status = request.form.get("from_status") or order.status

    if from_status not in ORDER_STATUSES or from_status == ORDER_STATUSES[-1]:
        return redirect(back_url)

    next_status = ORDER_STATUSES[ORDER_STATUSES.index(from_status) + 1]
    result = db.session.execute(
        Order.__table__.update()
        .where(Order.id == order.id, Order.status == from_status)
        .values(status=next_status)
    )
    db.session.commit()

    if result.rowcount == 0:
        flash(_("Status pesanan %(label)s sudah diubah dari device lain.", label=order.display_label), "warning")

    return redirect(back_url)


# ============================================================
# KASIR - pembayaran
# ============================================================

@staff_bp.route("/cashier")
@roles_required(ROLE_OWNER, ROLE_KASIR)
def cashier():
    orders = (
        Order.query.filter_by(is_paid=False)
        .order_by(Order.created_at.asc())
        .all()
    )

    today = datetime.now().date()
    start = datetime.combine(today, time.min)
    end = datetime.combine(today, time.max)

    paid_today = (
        Order.query.filter(
            Order.is_paid.is_(True),
            Order.paid_at >= start,
            Order.paid_at <= end,
        )
        .order_by(Order.paid_at.desc())
        .all()
    )

    return render_template("staff/cashier.html", orders=orders, paid_today=paid_today)


@staff_bp.route("/cashier/status")
@roles_required(ROLE_OWNER, ROLE_KASIR)
def cashier_status():
    """Sama seperti kitchen_status() tapi untuk pesanan yang menunggu
    dibayar - dipakai buat bunyikan notifikasi suara di halaman Kasir.
    Jumlah baris item ikut disertakan di "tanda tangan" supaya tagihan
    yang berubah (nambah menu setelah kasir sempat lihat) ikut memicu
    notifikasi, bukan cuma pesanan yang benar-benar baru."""

    signatures = [
        f"{order.id}:{len(order.items)}"
        for order in Order.query.filter_by(is_paid=False).all()
    ]
    return {"order_ids": signatures}


@staff_bp.route("/cashier/list")
@roles_required(ROLE_OWNER, ROLE_KASIR)
def cashier_list():
    """Sama seperti kitchen_list() - refresh daftar pesanan Kasir lewat
    AJAX supaya izin autoplay audio browser tidak ke-reset tiap ada
    pesanan baru."""

    orders = (
        Order.query.filter_by(is_paid=False)
        .order_by(Order.created_at.asc())
        .all()
    )
    return render_template("staff/_cashier_orders.html", orders=orders)


@staff_bp.route("/orders/paid-status")
@roles_required(ROLE_OWNER, ROLE_KASIR)
def paid_order_status():
    """Dipoll global (base.html) buat notifikasi "pembayaran selesai" -
    Owner dengar ini dari halaman manapun dia berada (pantau penjualan
    real-time), dan Kasir dengar kalau kasir LAIN yang proses transaksi
    (kasir yang proses sendiri sudah dapat konfirmasi instan langsung di
    halaman struk - lihat receipt(), tidak nunggu siklus poll ini)."""

    today = datetime.now().date()
    start = datetime.combine(today, time.min)
    end = datetime.combine(today, time.max)

    signatures = [
        f"{order.id}:{order.paid_at.isoformat()}"
        for order in Order.query.filter(
            Order.is_paid.is_(True),
            Order.paid_at >= start,
            Order.paid_at <= end,
        ).all()
    ]
    return {"order_ids": signatures}


@staff_bp.route("/orders/<int:order_id>/cancel", methods=["POST"])
@roles_required(ROLE_OWNER, ROLE_KASIR)
def cancel_order(order_id):
    """Batalkan pesanan yang belum lunas (salah input, tamu batal, dst)
    dan kembalikan stok bahan baku yang sudah kadung terpotong untuknya.
    Order yang SUDAH lunas sengaja tidak bisa dibatalkan lewat sini -
    pembatalan pesanan yang sudah dibayar butuh alur refund tersendiri
    di luar cakupan tombol ini."""

    order = Order.query.get_or_404(order_id)

    if order.is_paid:
        flash(_("Pesanan #%(id)s sudah dibayar, tidak bisa dibatalkan lewat sini.", id=order.id), "danger")
        return redirect(url_for("staff.cashier"))

    # Stok cuma dikembalikan kalau dapur BELUM mulai mengerjakan pesanan
    # ini - kalau sudah diproses/siap/diantar, bahan bakunya memang sudah
    # terpakai beneran, mengembalikannya bikin stok di sistem lebih
    # banyak dari stok fisik.
    restore_stock = order.status == "pending"
    if restore_stock:
        restore_stock_for_order(order)

    # Jejak pembatalan (lihat CancelledOrderLog) - item dibaca ulang dari
    # database, bukan dari objek di session, supaya tambahan tamu yang
    # baru saja masuk ikut tercatat.
    current_items = OrderItem.query.filter_by(order_id=order.id).all()
    db.session.add(CancelledOrderLog(
        order_id=order.id,
        label=f"{order.display_label} ({order.display_sublabel})",
        order_type=order.order_type,
        items_summary=", ".join(f"{item.quantity}x {item.name_snapshot}" for item in current_items) or "-",
        total=sum(item.subtotal for item in current_items),
        status_at_cancel=order.status,
        stock_restored=restore_stock,
        order_created_at=order.created_at,
        cancelled_by=current_user.username,
    ))

    # DELETE ... WHERE is_paid = 0 (bukan cek order.is_paid di atas yang
    # sudah basi) - jaga-jaga kasir lain baru saja memproses bayar
    # pesanan ini persis di antara cek tadi dan baris ini. Item & catatan
    # stok dihapus lewat subquery ke order yang sama di transaksi yang
    # sama, supaya item yang barusan ditambah tamu juga ikut terhapus.
    unpaid_order = select(Order.id).where(Order.id == order.id, Order.is_paid.is_(False))
    db.session.execute(OrderItem.__table__.delete().where(OrderItem.order_id.in_(unpaid_order)))
    db.session.execute(OrderStockDeduction.__table__.delete().where(OrderStockDeduction.order_id.in_(unpaid_order)))
    result = db.session.execute(
        Order.__table__.delete().where(Order.id == order.id, Order.is_paid.is_(False))
    )

    if result.rowcount == 0:
        db.session.rollback()
        flash(_("Pesanan #%(id)s sudah dibayar, tidak bisa dibatalkan lewat sini.", id=order_id), "danger")
        return redirect(url_for("staff.cashier"))

    db.session.commit()

    if restore_stock:
        flash(_("Pesanan #%(id)s dibatalkan, stok bahan baku yang terpakai sudah dikembalikan.", id=order_id), "success")
    else:
        flash(_("Pesanan #%(id)s dibatalkan. Stok tidak dikembalikan karena pesanan sudah mulai diproses dapur.", id=order_id), "success")
    return redirect(url_for("staff.cashier"))


@staff_bp.route("/orders/<int:order_id>/pay", methods=["POST"])
@roles_required(ROLE_OWNER, ROLE_KASIR)
def pay_order(order_id):
    order = Order.query.get_or_404(order_id)

    if order.is_paid:
        # Jalan cepat - hindarkan tamu/kasir dari isi form bayar buat
        # pesanan yang jelas-jelas sudah lunas. Penjaga SUNGGUHAN ada di
        # UPDATE ... WHERE is_paid = 0 di bawah (atomic di level DB),
        # bukan cek ini - cek Python biasa ini masih race-able kalau 2
        # submit datang nyaris bersamaan.
        flash(_("Pesanan #%(id)s sudah dibayar sebelumnya.", id=order.id), "warning")
        return redirect(url_for("staff.receipt", order_id=order.id))

    method = request.form.get("payment_method")

    if method not in PAYMENT_METHODS:
        flash(_("Metode pembayaran tidak valid."), "danger")
        return redirect(url_for("staff.cashier"))

    # Total yang TERLIHAT kasir di layar waktu tombol Bayar ditekan. Kalau
    # beda dengan total sekarang, berarti tamu baru saja nambah menu (atau
    # PPN baru diubah) setelah halaman Kasir dimuat - tolak supaya kasir
    # tidak menagih/menghitung kembalian pakai angka lama.
    expected_total = request.form.get("expected_total", type=int)
    if expected_total is not None and expected_total != order.grand_total:
        flash(_("Tagihan pesanan %(label)s baru saja berubah. Cek lagi totalnya sebelum bayar.", label=order.display_label), "warning")
        return redirect(url_for("staff.cashier"))

    settings = get_settings()
    if settings.ppn_enabled and settings.ppn_percentage:
        ppn_percentage = settings.ppn_percentage
        ppn_amount = calculate_ppn(order.total, settings.ppn_percentage)
    else:
        ppn_percentage = None
        ppn_amount = None

    grand_total = order.total + (ppn_amount or 0)

    cash_received = None
    change_amount = None

    if method == "cash":
        cash_received = request.form.get("cash_received", type=digits_int)

        if not cash_received or cash_received < grand_total:
            flash(_("Uang diterima kurang dari total tagihan."), "danger")
            return redirect(url_for("staff.cashier"))

        change_amount = cash_received - grand_total

    # UPDATE ... WHERE is_paid = 0 (bukan baca-lalu-tulis lewat ORM) -
    # supaya 2 submit pembayaran yang nyaris bersamaan untuk order yang
    # sama (klik 2x, 2 kasir buka order yang sama) tidak bisa berdua
    # lolos dan saling timpa payment_method/cash_received/served_by.
    # Submit kedua akan dapat rowcount 0 dan ditolak dengan pesan
    # "sudah dibayar", sama seperti pola di app/inventory.py.
    #
    # Syarat tambahan: total item di database masih sama persis dengan
    # yang dipakai menghitung PPN/kembalian di atas - kalau tamu keburu
    # nambah menu di antara hitungan tadi dan UPDATE ini, pembayaran
    # ditolak (bukan tercatat lunas dengan total & kembalian yang salah).
    current_items_total = (
        select(func.coalesce(func.sum(OrderItem.price_snapshot * OrderItem.quantity), 0))
        .where(OrderItem.order_id == order.id)
        .scalar_subquery()
    )
    result = db.session.execute(
        Order.__table__.update()
        .where(
            Order.id == order.id,
            Order.is_paid.is_(False),
            current_items_total == order.total,
        )
        .values(
            is_paid=True,
            payment_method=method,
            paid_at=datetime.now(),
            served_by=current_user.username,
            cash_received=cash_received,
            change_amount=change_amount,
            ppn_percentage=ppn_percentage,
            ppn_amount=ppn_amount,
        )
    )

    if result.rowcount == 0:
        db.session.rollback()
        db.session.expire_all()
        fresh = Order.query.get(order_id)
        if fresh is None:
            flash(_("Pesanan #%(id)s sudah dibatalkan.", id=order_id), "warning")
            return redirect(url_for("staff.cashier"))
        if not fresh.is_paid:
            flash(_("Tagihan pesanan %(label)s baru saja berubah. Cek lagi totalnya sebelum bayar.", label=fresh.display_label), "warning")
            return redirect(url_for("staff.cashier"))
        flash(_("Pesanan #%(id)s sudah dibayar sebelumnya.", id=order_id), "warning")
        return redirect(url_for("staff.receipt", order_id=order_id))

    db.session.commit()

    flash(_("Pesanan #%(id)s ditandai sudah dibayar.", id=order.id), "success")
    return redirect(url_for("staff.receipt", order_id=order.id, just_paid=1))


@staff_bp.route("/orders/<int:order_id>/receipt")
@roles_required(ROLE_OWNER, ROLE_KASIR)
def receipt(order_id):
    order = Order.query.get_or_404(order_id)
    if not order.is_paid:
        flash(_("Pesanan #%(id)s belum dibayar - struk baru tersedia setelah lunas.", id=order.id), "warning")
        return redirect(url_for("staff.cashier"))
    settings = get_settings()
    just_paid = request.args.get("just_paid") == "1"

    return render_template(
        "staff/receipt.html", order=order, settings=settings, just_paid=just_paid
    )


@staff_bp.route("/orders/<int:order_id>/receipt.pdf")
@roles_required(ROLE_OWNER, ROLE_KASIR)
def receipt_pdf(order_id):
    """Versi PDF dari struk (tombol "Cetak ke PDF") - dibuat manual pakai
    reportlab canvas (bukan platypus) karena lebar kertasnya sempit ala
    printer thermal (58/80mm) dan tingginya harus mengikuti panjang isi
    struk, bukan ukuran kertas baku seperti A4."""
    import io

    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdf_canvas

    order = Order.query.get_or_404(order_id)
    if not order.is_paid:
        abort(404)
    settings = get_settings()

    page_w = (58 if settings.receipt_paper_width == "58" else 80) * mm
    margin = 4 * mm
    content_w = page_w - 2 * margin
    line_h = 4.6 * mm
    font_size = 8 if settings.receipt_paper_width == "58" else 9

    logo_filename = (
        order.channel.logo if (order.channel and order.channel.logo)
        else settings.receipt_logo_filename
    )
    logo_reader = None
    logo_h = 0
    if logo_filename:
        logo_path = os.path.join(current_app.static_folder, "uploads", "branding", logo_filename)
        if os.path.exists(logo_path):
            logo_reader = ImageReader(logo_path)
            iw, ih = logo_reader.getSize()
            logo_h = min(18 * mm, content_w * 0.5 * ih / iw)

    # Kumpulkan baris dulu (label, value, style) supaya tinggi halaman bisa
    # dihitung SEBELUM canvas dibuat - reportlab butuh ukuran halaman di awal.
    header_lines = []
    if order.channel and order.channel.logo:
        header_lines.append(("center-bold", order.channel.name))
        header_lines.append(("center-muted", settings.shop_name))
    else:
        header_lines.append(("center-bold", settings.shop_name))
        if settings.address:
            header_lines.append(("center", settings.address))
        if settings.phone:
            header_lines.append(("center", _("Telp: %(phone)s", phone=settings.phone)))
    header_lines.append(("center-muted", _("*** STRUK PEMBAYARAN ***")))

    info_rows = [
        (_("No. Struk"), f"#{order.id}"),
        (_("Meja") if order.table else _("Jenis"), f"{order.display_label} ({order.display_sublabel})"),
        (_("Tanggal"), (order.paid_at or order.created_at).strftime("%d/%m/%Y %H:%M")),
    ]
    if order.served_by:
        info_rows.append((_("Kasir"), order.served_by))

    item_lines = []
    for item in order.items:
        item_lines.append(("left", item.name_snapshot))
        item_lines.append((
            "row-muted",
            f"{item.quantity} x {'{:,}'.format(item.price_snapshot).replace(',', '.')}",
            "{:,}".format(item.subtotal).replace(",", "."),
        ))

    ppn_rows = []
    if order.ppn_amount:
        ppn_rows.append((_("Subtotal"), f"Rp {'{:,}'.format(order.total).replace(',', '.')}"))
        ppn_rows.append((
            _("PPN %(pct)s%%", pct=round(order.ppn_percentage, 1)),
            f"Rp {'{:,}'.format(order.ppn_amount).replace(',', '.')}",
        ))

    payment_rows = [(_("Bayar"), "Cash" if order.payment_method == "cash" else "QRIS")]
    if order.payment_method == "cash" and order.cash_received is not None:
        payment_rows.append((_("Tunai"), f"Rp {'{:,}'.format(order.cash_received).replace(',', '.')}"))
        payment_rows.append((_("Kembalian"), f"Rp {'{:,}'.format(order.change_amount).replace(',', '.')}"))

    social_lines = []
    if settings.instagram:
        social_lines.append(_("IG: %(handle)s", handle=settings.instagram))
    if settings.tiktok:
        social_lines.append(_("TikTok: %(handle)s", handle=settings.tiktok))
    if settings.whatsapp:
        social_lines.append(_("WA: %(handle)s", handle=settings.whatsapp))
    if settings.other_social:
        social_lines.append(settings.other_social)

    # Hitung tinggi total: tiap "blok" di atas + jarak antar-divider.
    n_lines = (
        len(header_lines) + len(info_rows) + len(item_lines)
        + len(ppn_rows) + 1 + len(payment_rows) + len(social_lines)
        + 4  # "Orulabs (c) 2026", "Terima kasih!", dan 2 baris jarak ekstra
    )
    n_dividers = 3 + (1 if ppn_rows else 0) + (1 if social_lines else 0)
    page_h = (
        margin * 2 + logo_h + (logo_h and 2 * mm)
        + n_lines * line_h + n_dividers * 2.5 * mm + 4 * mm
    )

    buffer = io.BytesIO()
    c = pdf_canvas.Canvas(buffer, pagesize=(page_w, page_h))
    y = page_h - margin

    def center(text, bold=False, muted=False, size=None):
        nonlocal y
        c.setFont("Courier-Bold" if bold else "Courier", size or font_size)
        c.setFillGray(0.33 if muted else 0)
        c.drawCentredString(page_w / 2, y - font_size * 0.8, _pdf_safe(text))
        y -= line_h

    def row(left, right, bold=False, muted=False):
        nonlocal y
        c.setFont("Courier-Bold" if bold else "Courier", font_size)
        c.setFillGray(0.27 if muted else 0)
        c.drawString(margin, y - font_size * 0.8, _pdf_safe(left))
        c.drawRightString(page_w - margin, y - font_size * 0.8, _pdf_safe(right))
        y -= line_h

    def divider(dashed=True):
        nonlocal y
        y -= 1 * mm
        c.setDash([1.5, 1.5] if dashed else [])
        c.setLineWidth(0.6)
        c.line(margin, y, page_w - margin, y)
        c.setDash([])
        y -= 2.5 * mm

    if logo_reader:
        logo_w = logo_h * (logo_reader.getSize()[0] / logo_reader.getSize()[1])
        c.drawImage(
            logo_reader, (page_w - logo_w) / 2, y - logo_h,
            width=logo_w, height=logo_h, mask="auto",
        )
        y -= logo_h + 2 * mm

    for kind, text in header_lines:
        if kind == "center-bold":
            center(text, bold=True)
        elif kind == "center-muted":
            center(text, muted=True)
        else:
            center(text)

    divider()
    for label, value in info_rows:
        row(label, value)
    divider()

    for entry in item_lines:
        if entry[0] == "left":
            row(entry[1], "")
        else:
            row(entry[1], entry[2], muted=True)

    if ppn_rows:
        divider()
        for label, value in ppn_rows:
            row(label, value)

    divider(dashed=False)
    row(_("TOTAL"), f"Rp {'{:,}'.format(order.grand_total).replace(',', '.')}", bold=True)
    divider(dashed=False)

    for label, value in payment_rows:
        row(label, value, muted=True)

    if social_lines:
        divider()
        for line in social_lines:
            center(line)

    center("Orulabs © 2026", muted=True, size=max(6, font_size - 1))
    divider(dashed=False)
    center(_("Terima kasih!"), bold=True)

    c.showPage()
    c.save()
    buffer.seek(0)

    filename = f"struk-{order.id}.pdf"
    return send_file(
        buffer, mimetype="application/pdf",
        as_attachment=True, download_name=filename,
    )


# ============================================================
# ADMIN - KATEGORI & MENU
# ============================================================

@staff_bp.route("/admin/menu", methods=["GET", "POST"])
@roles_required(ROLE_OWNER)
def admin_menu():
    if request.method == "POST":
        form_type = request.form.get("form_type")

        if form_type == "category":
            name = request.form.get("name", "").strip()

            if name:
                db.session.add(Category(name=name))
                db.session.commit()
                flash(_("Kategori ditambahkan."), "success")

        elif form_type == "menu_item":
            name = request.form.get("name", "").strip()
            price = request.form.get("price", type=digits_int)
            category_id = request.form.get("category_id", type=int)
            category = Category.query.get(category_id) if category_id else None

            if name and price is not None and price > 0 and category:
                db.session.add(
                    MenuItem(name=name, price=price, category_id=category.id)
                )
                db.session.commit()
                flash(_("Menu ditambahkan."), "success")
            else:
                flash(_("Isi nama, harga (lebih dari 0), dan kategori menu dengan benar."), "warning")

        return redirect(url_for("staff.admin_menu"))

    categories = Category.query.order_by(Category.order, Category.name).all()
    ingredients = Ingredient.query.order_by(Ingredient.name).all()
    return render_template(
        "staff/menu_admin.html", categories=categories, ingredients=ingredients
    )


@staff_bp.route("/admin/menu/<int:item_id>/toggle", methods=["POST"])
@roles_required(ROLE_OWNER)
def toggle_menu_item(item_id):
    item = MenuItem.query.get_or_404(item_id)
    item.is_available = not item.is_available
    db.session.commit()
    return redirect(url_for("staff.admin_menu"))


@staff_bp.route("/admin/menu/<int:item_id>/delete", methods=["POST"])
@roles_required(ROLE_OWNER)
def delete_menu_item(item_id):
    item = MenuItem.query.get_or_404(item_id)

    if item.order_items:
        flash(
            _(
                "Menu %(name)s tidak bisa dihapus karena sudah punya riwayat "
                "pesanan (dipakai di laporan/struk). Tandai \"Habis\" saja kalau "
                "sudah tidak dijual.",
                name=item.name,
            ),
            "danger",
        )
        return redirect(url_for("staff.admin_menu"))

    _remove_logo(item, "photo", subfolder="menu")

    # Harga khusus platform ojol untuk menu ini ikut dihapus - kalau
    # dibiarkan, menu BARU yang kebetulan dapat id yang sama akan diam-diam
    # mewarisi harga lama itu.
    MenuItemChannelPrice.query.filter_by(menu_item_id=item.id).delete()

    name = item.name
    db.session.delete(item)
    db.session.commit()
    flash(_("Menu %(name)s dihapus.", name=name), "success")
    return redirect(url_for("staff.admin_menu"))


@staff_bp.route("/admin/menu/<int:item_id>/photo", methods=["POST"])
@roles_required(ROLE_OWNER)
def update_menu_item_photo(item_id):
    item = MenuItem.query.get_or_404(item_id)

    if request.form.get("remove_photo") == "1":
        _remove_logo(item, "photo", subfolder="menu")
        db.session.commit()
        flash(_("Foto %(name)s dihapus.", name=item.name), "success")
        return redirect(url_for("staff.admin_menu"))

    error = _save_logo(
        "photo", item, "photo",
        subfolder="menu", filename_base=f"menu-{item.id}",
    )

    if error:
        flash(error, "danger")
        return redirect(url_for("staff.admin_menu"))

    db.session.commit()
    flash(_("Foto %(name)s berhasil diperbarui.", name=item.name), "success")
    return redirect(url_for("staff.admin_menu"))


@staff_bp.route("/admin/menu/<int:item_id>/recipe/add", methods=["POST"])
@roles_required(ROLE_OWNER)
def add_recipe_item(item_id):
    item = MenuItem.query.get_or_404(item_id)
    ingredient_id = request.form.get("ingredient_id", type=int)
    quantity_used = request.form.get("quantity_used", type=finite_float)

    ingredient = Ingredient.query.get(ingredient_id) if ingredient_id else None

    if not ingredient or not quantity_used or quantity_used <= 0:
        flash(_("Pilih bahan baku dan isi jumlah pemakaian dengan benar."), "warning")
        return redirect(url_for("staff.admin_menu"))

    existing = MenuItemIngredient.query.filter_by(
        menu_item_id=item.id, ingredient_id=ingredient.id
    ).first()

    if existing:
        existing.quantity_used = quantity_used
    else:
        db.session.add(
            MenuItemIngredient(
                menu_item_id=item.id,
                ingredient_id=ingredient.id,
                quantity_used=quantity_used,
            )
        )

    db.session.commit()
    flash(_("Resep %(name)s berhasil diperbarui.", name=item.name), "success")
    return redirect(url_for("staff.admin_menu"))


@staff_bp.route("/admin/menu/<int:item_id>/recipe/<int:ingredient_id>/remove", methods=["POST"])
@roles_required(ROLE_OWNER)
def remove_recipe_item(item_id, ingredient_id):
    row = MenuItemIngredient.query.filter_by(
        menu_item_id=item_id, ingredient_id=ingredient_id
    ).first_or_404()
    db.session.delete(row)
    db.session.commit()
    return redirect(url_for("staff.admin_menu"))


# ============================================================
# ADMIN - INVENTORY BAHAN BAKU
# ============================================================

@staff_bp.route("/admin/inventory", methods=["GET", "POST"])
@roles_required(ROLE_OWNER)
def admin_inventory():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        unit = request.form.get("unit", "").strip()
        stock_quantity = request.form.get("stock_quantity", type=finite_float)
        low_stock_threshold = request.form.get("low_stock_threshold", type=finite_float)

        if name and unit and stock_quantity is not None and stock_quantity >= 0 and (low_stock_threshold or 0) >= 0:
            db.session.add(
                Ingredient(
                    name=name,
                    unit=unit,
                    stock_quantity=stock_quantity,
                    low_stock_threshold=low_stock_threshold or 0,
                )
            )
            db.session.commit()
            flash(_("Bahan baku %(name)s ditambahkan.", name=name), "success")
        else:
            flash(_("Isi nama, satuan, dan stok awal dengan benar."), "warning")

        return redirect(url_for("staff.admin_inventory"))

    ingredients = Ingredient.query.order_by(Ingredient.name).all()
    return render_template("staff/inventory.html", ingredients=ingredients)


@staff_bp.route("/admin/inventory/<int:ingredient_id>/restock", methods=["POST"])
@roles_required(ROLE_OWNER)
def restock_ingredient(ingredient_id):
    ingredient = Ingredient.query.get_or_404(ingredient_id)
    amount = request.form.get("amount", type=finite_float)

    if not amount or amount <= 0:
        flash(_("Jumlah restock tidak valid."), "warning")
        return redirect(url_for("staff.admin_inventory"))

    # UPDATE ... SET stock_quantity = stock_quantity + amount (bukan
    # baca-lalu-tulis lewat ORM) - dua restock yang nyaris bersamaan
    # (2 device, atau submit dobel) tetap terjumlah dengan benar, tidak
    # ada yang ketiban timpa. Pola sama seperti pengurangan stok di
    # app/inventory.py.
    db.session.execute(
        Ingredient.__table__.update()
        .where(Ingredient.id == ingredient.id)
        .values(stock_quantity=Ingredient.stock_quantity + amount)
    )
    db.session.commit()
    db.session.refresh(ingredient)
    flash(
        _("Stok %(name)s ditambah %(amount)s %(unit)s.", name=ingredient.name, amount=amount, unit=ingredient.unit),
        "success",
    )
    return redirect(url_for("staff.admin_inventory"))


@staff_bp.route("/admin/inventory/<int:ingredient_id>/delete", methods=["POST"])
@roles_required(ROLE_OWNER)
def delete_ingredient(ingredient_id):
    ingredient = Ingredient.query.get_or_404(ingredient_id)

    if ingredient.recipe_usages:
        menu_names = ", ".join(sorted({row.menu_item.name for row in ingredient.recipe_usages}))
        flash(
            _(
                "Bahan %(name)s tidak bisa dihapus karena masih dipakai di resep: %(menus)s.",
                name=ingredient.name,
                menus=menu_names,
            ),
            "danger",
        )
        return redirect(url_for("staff.admin_inventory"))

    name = ingredient.name
    db.session.delete(ingredient)
    db.session.commit()
    flash(_("Bahan %(name)s dihapus.", name=name), "success")
    return redirect(url_for("staff.admin_inventory"))


INVENTORY_EXPORT_SCOPES = ("full", "low-stock")


def _inventory_export_ingredients(scope):
    ingredients = Ingredient.query.order_by(Ingredient.name).all()
    if scope == "low-stock":
        ingredients = [i for i in ingredients if i.is_low_stock]
    return ingredients


@staff_bp.route("/admin/inventory/export/<scope>.xlsx")
@roles_required(ROLE_OWNER)
def export_inventory_excel(scope):
    if scope not in INVENTORY_EXPORT_SCOPES:
        abort(404)

    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    ingredients = _inventory_export_ingredients(scope)
    settings = get_settings()
    is_low_stock = scope == "low-stock"
    title = _("Bahan Baku Menipis / Habis") if is_low_stock else _("Daftar Bahan Baku")

    wb = Workbook()
    ws = wb.active
    # Nama sheet Excel tidak boleh mengandung \ / ? * [ ] : - beda dari
    # judul yang ditampilkan di isi filenya (title di atas, boleh bebas).
    sheet_title = str(title)
    for ch in r"\/?*[]:":
        sheet_title = sheet_title.replace(ch, "-")
    ws.title = sheet_title[:31]

    header_fill = PatternFill("solid", fgColor="0E4C75")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=14)
    danger_font = Font(color="DC2626", bold=True)
    thin_border = Border(bottom=Side(style="thin", color="E2E8F0"))

    ws.merge_cells("A1:E1")
    ws["A1"] = _excel_safe(settings.shop_name)
    ws["A1"].font = title_font

    ws.merge_cells("A2:E2")
    ws["A2"] = f"{title} - {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(color="64748B")

    header_row = 4
    headers = [_("Nama Bahan"), _("Satuan"), _("Stok Saat Ini"), _("Batas Stok Menipis"), _("Status")]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=str(text))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    row = header_row + 1
    for ingredient in ingredients:
        values = [
            ingredient.name,
            ingredient.unit,
            ingredient.stock_quantity,
            ingredient.low_stock_threshold,
            str(_("Menipis")) if ingredient.is_low_stock else str(_("Aman")),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=_excel_safe(value))
            cell.border = thin_border
            if col == 5 and ingredient.is_low_stock:
                cell.font = danger_font
        row += 1

    if not ingredients:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        empty_text = _("Tidak ada bahan baku yang menipis/habis.") if is_low_stock else _("Belum ada bahan baku.")
        ws.cell(row=row, column=1, value=str(empty_text))
        row += 1

    ws.freeze_panes = f"A{header_row + 1}"

    widths = [26, 12, 16, 20, 12]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"inventory-{scope}-{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True, download_name=filename,
    )


@staff_bp.route("/admin/inventory/export/<scope>.pdf")
@roles_required(ROLE_OWNER)
def export_inventory_pdf(scope):
    if scope not in INVENTORY_EXPORT_SCOPES:
        abort(404)

    import io

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    ingredients = _inventory_export_ingredients(scope)
    settings = get_settings()
    is_low_stock = scope == "low-stock"
    title = _("Bahan Baku Menipis / Habis") if is_low_stock else _("Daftar Bahan Baku")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        topMargin=16 * mm, bottomMargin=16 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm,
    )
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(xml_escape(settings.shop_name), styles["Title"]))
    story.append(Paragraph(f"{title} &middot; {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 8 * mm))

    header = [_("Nama Bahan"), _("Satuan"), _("Stok Saat Ini"), _("Batas Stok Menipis"), _("Status")]
    table_data = [header]

    for ingredient in ingredients:
        status_text = str(_("Menipis")) if ingredient.is_low_stock else str(_("Aman"))
        table_data.append([
            ingredient.name,
            ingredient.unit,
            "{:g}".format(ingredient.stock_quantity),
            "{:g}".format(ingredient.low_stock_threshold),
            status_text,
        ])

    if len(table_data) == 1:
        empty_text = _("Tidak ada bahan baku yang menipis/habis.") if is_low_stock else _("Belum ada bahan baku.")
        table_data.append([str(empty_text), "", "", "", ""])

    ingredients_table = Table(
        table_data,
        colWidths=[55 * mm, 25 * mm, 30 * mm, 35 * mm, 25 * mm],
        repeatRows=1,
    )
    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E4C75")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F9FC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for i, ingredient in enumerate(ingredients, start=1):
        if ingredient.is_low_stock:
            style_commands.append(("TEXTCOLOR", (4, i), (4, i), colors.HexColor("#DC2626")))
            style_commands.append(("FONTNAME", (4, i), (4, i), "Helvetica-Bold"))
    ingredients_table.setStyle(TableStyle(style_commands))
    story.append(ingredients_table)

    doc.build(story)
    buffer.seek(0)

    filename = f"inventory-{scope}-{datetime.now().strftime('%Y%m%d')}.pdf"
    return send_file(
        buffer, mimetype="application/pdf",
        as_attachment=True, download_name=filename,
    )


# ============================================================
# HELPER - QR CODE MEJA (dipakai oleh table_map())
# ============================================================

def _generate_table_qr(table):
    """Generate gambar QR code yang mengarah ke halaman menu meja ini,
    disimpan di app/static/qrcodes/<code>.png supaya bisa dicetak/tempel
    di meja fisik."""

    target_url = f"{current_app.config['BASE_URL']}/t/{table.code}"

    qr_dir = os.path.join(current_app.root_path, "static", "qrcodes")
    os.makedirs(qr_dir, exist_ok=True)

    img = qrcode.make(target_url)
    img.save(os.path.join(qr_dir, f"{table.code}.png"))


# ============================================================
# LAPORAN SEDERHANA
# ============================================================

def _period_report(start, end):
    """Rekap transaksi lunas dalam satu rentang waktu (dipakai untuk
    rekap harian/mingguan/bulanan/tahunan di halaman Laporan)."""

    orders = Order.query.options(selectinload(Order.items)).filter(
        Order.is_paid.is_(True),
        Order.paid_at >= start,
        Order.paid_at <= end,
    ).all()

    total_sales = sum(order.total for order in orders)
    ppn_collected = sum(order.ppn_amount or 0 for order in orders)
    transaction_count = len(orders)

    item_counts = {}
    items_sold = 0

    for order in orders:
        for item in order.items:
            item_counts[item.name_snapshot] = (
                item_counts.get(item.name_snapshot, 0) + item.quantity
            )
            items_sold += item.quantity

    top_items = sorted(item_counts.items(), key=lambda pair: pair[1], reverse=True)[:10]

    cancellations = (
        CancelledOrderLog.query.filter(
            CancelledOrderLog.cancelled_at >= start,
            CancelledOrderLog.cancelled_at <= end,
        )
        .order_by(CancelledOrderLog.cancelled_at.desc())
        .all()
    )

    return {
        "total_sales": total_sales,
        "ppn_collected": ppn_collected,
        "transaction_count": transaction_count,
        "items_sold": items_sold,
        "avg_per_transaction": (total_sales / transaction_count) if transaction_count else 0,
        "top_items": top_items,
        "cancellations": cancellations,
        "cancelled_count": len(cancellations),
        "cancelled_total": sum(log.total for log in cancellations),
    }


BULAN_ID = [
    _l("Januari"), _l("Februari"), _l("Maret"), _l("April"), _l("Mei"), _l("Juni"),
    _l("Juli"), _l("Agustus"), _l("September"), _l("Oktober"), _l("November"), _l("Desember"),
]
BULAN_ID_SHORT = [
    _l("Jan"), _l("Feb"), _l("Mar"), _l("Apr"), _l("Mei"), _l("Jun"),
    _l("Jul"), _l("Agu"), _l("Sep"), _l("Okt"), _l("Nov"), _l("Des"),
]

PERIOD_KEYS = ("daily", "weekly", "monthly", "yearly")
PERIOD_LABELS = {
    "daily": _l("Harian"),
    "weekly": _l("Mingguan"),
    "monthly": _l("Bulanan"),
    "yearly": _l("Tahunan"),
}


def _period_bounds(period_key):
    """Hitung (start, end, range_label) untuk satu period_key ('daily',
    'weekly', 'monthly', 'yearly'), dipakai bareng oleh reports() dan
    route export CSV/PDF supaya rentang tanggalnya selalu sinkron."""

    today = datetime.now().date()

    if period_key == "daily":
        start_date, end_date = today, today
        range_label = f"{today.day} {BULAN_ID[today.month - 1]} {today.year}"

    elif period_key == "weekly":
        start_date = today - timedelta(days=today.weekday())
        end_date = start_date + timedelta(days=6)
        range_label = (
            f"{start_date.day} {BULAN_ID_SHORT[start_date.month - 1]} - "
            f"{end_date.day} {BULAN_ID_SHORT[end_date.month - 1]} {end_date.year}"
        )

    elif period_key == "monthly":
        start_date = today.replace(day=1)
        next_month_date = (
            start_date.replace(year=start_date.year + 1, month=1)
            if start_date.month == 12
            else start_date.replace(month=start_date.month + 1)
        )
        end_date = next_month_date - timedelta(days=1)
        range_label = f"{BULAN_ID[today.month - 1]} {today.year}"

    elif period_key == "yearly":
        start_date = today.replace(month=1, day=1)
        end_date = today.replace(month=12, day=31)
        range_label = str(today.year)

    else:
        abort(404)

    return (
        datetime.combine(start_date, time.min),
        datetime.combine(end_date, time.max),
        range_label,
    )


@staff_bp.route("/reports")
@roles_required(ROLE_OWNER)
def reports():
    periods = {}

    for key in PERIOD_KEYS:
        start, end, range_label = _period_bounds(key)
        periods[key] = {
            "label": PERIOD_LABELS[key],
            "range_label": range_label,
            "data": _period_report(start, end),
        }

    return render_template("staff/reports.html", periods=periods)


def _orders_in_period(start, end):
    return (
        Order.query.options(selectinload(Order.items)).filter(
            Order.is_paid.is_(True),
            Order.paid_at >= start,
            Order.paid_at <= end,
        )
        .order_by(Order.paid_at)
        .all()
    )


@staff_bp.route("/reports/export/<period_key>.xlsx")
@roles_required(ROLE_OWNER)
def export_report_excel(period_key):
    if period_key not in PERIOD_KEYS:
        abort(404)

    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    start, end, range_label = _period_bounds(period_key)
    orders = _orders_in_period(start, end)
    data = _period_report(start, end)
    settings = get_settings()

    wb = Workbook()
    ws = wb.active
    ws.title = str(PERIOD_LABELS[period_key])[:31]

    header_fill = PatternFill("solid", fgColor="5C3B1E")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=14)
    bold_font = Font(bold=True)
    thin_border = Border(bottom=Side(style="thin", color="E4D9CC"))
    rupiah_format = '"Rp" #,##0'

    ws.merge_cells("A1:J1")
    ws["A1"] = settings.shop_name
    ws["A1"].font = title_font

    ws.merge_cells("A2:J2")
    ws["A2"] = f"{PERIOD_LABELS[period_key]} - {range_label}"
    ws["A2"].font = Font(color="6B5B4C")

    summary_rows = [
        (_("Total Penjualan"), data["total_sales"], rupiah_format),
        (_("PPN Terkumpul"), data["ppn_collected"], rupiah_format),
        (_("Transaksi"), data["transaction_count"], "0"),
        (_("Rata-rata/Transaksi"), round(data["avg_per_transaction"]), rupiah_format),
        (_("Item Terjual"), data["items_sold"], "0"),
    ]
    row = 4
    for label, value, num_format in summary_rows:
        ws.cell(row=row, column=1, value=str(label)).font = bold_font
        cell = ws.cell(row=row, column=2, value=value)
        cell.number_format = num_format
        cell.font = bold_font
        row += 1

    row += 1
    header_row = row
    headers = [
        _("No. Struk"), _("Tanggal"), _("Meja"), _("Item"),
        _("Metode"), _("PPN"), _("Total"), _("Tunai"), _("Kembalian"), _("Kasir"),
    ]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=str(text))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    row = header_row + 1
    total_sales = 0
    total_ppn = 0
    total_grand = 0

    for order in orders:
        items_desc = ", ".join(f"{i.quantity}x {i.name_snapshot}" for i in order.items)
        values = [
            f"#{order.id}",
            order.paid_at.strftime("%d/%m/%Y %H:%M"),
            order.display_label,
            items_desc,
            "Cash" if order.payment_method == "cash" else "QRIS",
            order.ppn_amount or 0,
            order.grand_total,
            order.cash_received,
            order.change_amount,
            order.served_by or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=_excel_safe(value))
            cell.border = thin_border
            if col in (6, 7, 8, 9) and value is not None:
                cell.number_format = rupiah_format
        total_sales += order.total
        total_ppn += order.ppn_amount or 0
        total_grand += order.grand_total
        row += 1

    if not orders:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.cell(row=row, column=1, value=str(_("Belum ada penjualan di periode ini.")))
        row += 1

    row += 1
    ws.cell(row=row, column=6, value=str(_("Total Penjualan"))).font = bold_font
    total_cell = ws.cell(row=row, column=7, value=total_sales)
    total_cell.number_format = rupiah_format
    total_cell.font = bold_font
    row += 1
    ws.cell(row=row, column=6, value=str(_("PPN Terkumpul"))).font = bold_font
    ppn_cell = ws.cell(row=row, column=7, value=total_ppn)
    ppn_cell.number_format = rupiah_format
    ppn_cell.font = bold_font
    row += 1
    ws.cell(row=row, column=6, value=str(_("Total Termasuk PPN"))).font = bold_font
    grand_cell = ws.cell(row=row, column=7, value=total_grand)
    grand_cell.number_format = rupiah_format
    grand_cell.font = bold_font
    row += 1
    ws.cell(row=row, column=6, value=str(_("Jumlah Transaksi"))).font = bold_font
    ws.cell(row=row, column=7, value=len(orders)).font = bold_font

    ws.freeze_panes = f"A{header_row + 1}"

    widths = [10, 17, 12, 40, 9, 12, 13, 13, 13, 14]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"laporan-{period_key}-{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True, download_name=filename,
    )


@staff_bp.route("/reports/export/<period_key>.pdf")
@roles_required(ROLE_OWNER)
def export_report_pdf(period_key):
    if period_key not in PERIOD_KEYS:
        abort(404)

    import io

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    start, end, range_label = _period_bounds(period_key)
    orders = _orders_in_period(start, end)
    data = _period_report(start, end)
    settings = get_settings()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=16 * mm, bottomMargin=16 * mm,
        leftMargin=14 * mm, rightMargin=14 * mm,
    )
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(xml_escape(settings.shop_name), styles["Title"]))
    story.append(Paragraph(f"{PERIOD_LABELS[period_key]} &middot; {range_label}", styles["Normal"]))
    story.append(Spacer(1, 10 * mm))

    summary_rows = [
        [_("Total Penjualan"), f"Rp {data['total_sales']:,}".replace(",", ".")],
        [_("PPN Terkumpul"), f"Rp {data['ppn_collected']:,}".replace(",", ".")],
        [_("Transaksi"), str(data["transaction_count"])],
        [_("Rata-rata/Transaksi"), f"Rp {data['avg_per_transaction']:,.0f}".replace(",", ".")],
        [_("Item Terjual"), str(data["items_sold"])],
    ]
    summary_table = Table(summary_rows, colWidths=[55 * mm, 45 * mm])
    summary_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E4D9CC")),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 8 * mm))

    header = [
        _("No. Struk"), _("Tanggal"), _("Meja"), _("Item"),
        _("Metode"), _("PPN"), _("Total"), _("Tunai"), _("Kembalian"), _("Kasir"),
    ]
    table_data = [header]

    for order in orders:
        items_desc = ", ".join(f"{i.quantity}x {i.name_snapshot}" for i in order.items)
        table_data.append([
            f"#{order.id}",
            order.paid_at.strftime("%d/%m/%Y %H:%M"),
            order.display_label,
            Paragraph(xml_escape(items_desc), styles["Normal"]),
            "Cash" if order.payment_method == "cash" else "QRIS",
            f"Rp {order.ppn_amount:,}".replace(",", ".") if order.ppn_amount else "-",
            f"Rp {order.grand_total:,}".replace(",", "."),
            f"Rp {order.cash_received:,}".replace(",", ".") if order.cash_received else "-",
            f"Rp {order.change_amount:,}".replace(",", ".") if order.change_amount else "-",
            order.served_by or "-",
        ])

    if len(table_data) == 1:
        table_data.append([_("Belum ada penjualan di periode ini."), "", "", "", "", "", "", "", "", ""])

    orders_table = Table(
        table_data,
        colWidths=[14 * mm, 24 * mm, 18 * mm, 52 * mm, 14 * mm, 20 * mm, 22 * mm, 20 * mm, 20 * mm, 20 * mm],
        repeatRows=1,
    )
    orders_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5C3B1E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E4D9CC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FBF3E8")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(orders_table)

    doc.build(story)
    buffer.seek(0)

    filename = f"laporan-{period_key}-{datetime.now().strftime('%Y%m%d')}.pdf"
    return send_file(
        buffer, mimetype="application/pdf",
        as_attachment=True, download_name=filename,
    )


# ============================================================
# ADMIN - USER & HAK AKSES
# ============================================================

@staff_bp.route("/admin/users", methods=["GET", "POST"])
@roles_required(ROLE_OWNER)
def admin_users():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role")

        if not username or not password or role not in ROLES:
            flash(_("Isi username, password, dan role dengan benar."), "danger")
        elif User.query.filter_by(username=username).first():
            flash(_("Username sudah dipakai."), "danger")
        else:
            user = User(username=username, role=role)
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash(
                _(
                    "User '%(username)s' (%(role)s) berhasil dibuat.",
                    username=username,
                    role=ROLE_LABELS[role],
                ),
                "success",
            )

        return redirect(url_for("staff.admin_users"))

    users = User.query.order_by(User.username).all()
    recent_logs = LoginLog.query.order_by(LoginLog.created_at.desc()).limit(50).all()

    return render_template(
        "staff/users_admin.html",
        users=users,
        roles=ROLES,
        role_labels=ROLE_LABELS,
        recent_logs=recent_logs,
    )


@staff_bp.route("/admin/users/status")
@roles_required(ROLE_OWNER)
def admin_users_status():
    """Dipoll berkala lewat JS di halaman Kelola User, supaya indikator
    online tiap akun terasa real-time tanpa perlu reload seluruh
    halaman (biar form Tambah User yang sedang diisi tidak keganggu)."""

    users = User.query.all()

    return {
        "users": [
            {
                "id": user.id,
                "is_online": user.is_online,
                "last_seen_label": (
                    user.last_seen_at.strftime("%d/%m %H:%M")
                    if user.last_seen_at else None
                ),
            }
            for user in users
        ]
    }


@staff_bp.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@roles_required(ROLE_OWNER)
def toggle_user(user_id):
    user = User.query.get_or_404(user_id)

    if user.id == current_user.id:
        flash(_("Tidak bisa menonaktifkan akun sendiri."), "warning")
        return redirect(url_for("staff.admin_users"))

    user.is_active_user = not user.is_active_user
    db.session.commit()

    return redirect(url_for("staff.admin_users"))


@staff_bp.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@roles_required(ROLE_OWNER)
def delete_user(user_id):
    user = User.query.get_or_404(user_id)

    if user.id == current_user.id:
        flash(_("Tidak bisa menghapus akun sendiri."), "warning")
        return redirect(url_for("staff.admin_users"))

    _remove_logo(user, "photo", subfolder="avatars")

    username = user.username
    db.session.delete(user)
    db.session.commit()
    flash(_("User %(username)s dihapus.", username=username), "success")
    return redirect(url_for("staff.admin_users"))


@staff_bp.route("/admin/users/<int:user_id>/photo", methods=["POST"])
@roles_required(ROLE_OWNER)
def update_user_photo(user_id):
    user = User.query.get_or_404(user_id)

    if request.form.get("remove_photo") == "1":
        _remove_logo(user, "photo", subfolder="avatars")
        db.session.commit()
        flash(_("Foto %(username)s dihapus.", username=user.username), "success")
        return redirect(url_for("staff.admin_users"))

    error = _save_logo(
        "photo", user, "photo",
        subfolder="avatars", filename_base=f"user-{user.id}",
    )

    if error:
        flash(error, "danger")
        return redirect(url_for("staff.admin_users"))

    db.session.commit()
    flash(_("Foto profil %(username)s berhasil diperbarui.", username=user.username), "success")
    return redirect(url_for("staff.admin_users"))


# ============================================================
# ADMIN - PENGATURAN TOKO (NAMA, LOGO, FOOTER)
# ============================================================

@staff_bp.route("/admin/settings", methods=["GET", "POST"])
@roles_required(ROLE_OWNER)
def admin_settings():
    settings = get_settings()

    if request.method == "POST":
        shop_name = request.form.get("shop_name", "").strip()

        if not shop_name:
            flash(_("Nama toko wajib diisi."), "danger")
            return redirect(url_for("staff.admin_settings"))

        # Nilai identitas toko SEBELUM form ini diterapkan - dipakai di
        # bawah buat tahu field mana yang benar-benar diubah Owner selama
        # Mode Demo aktif (lihat _sync_demo_backup()).
        values_before = {field: getattr(settings, field) for field in DEMO_BACKUP_FIELDS}

        error = _save_logo("logo_square", settings, "logo_square")
        if error:
            flash(error, "danger")
            return redirect(url_for("staff.admin_settings"))

        error = _save_logo("logo_wide", settings, "logo_wide")
        if error:
            flash(error, "danger")
            return redirect(url_for("staff.admin_settings"))

        error = _save_logo("qris_image", settings, "qris_image")
        if error:
            flash(error, "danger")
            return redirect(url_for("staff.admin_settings"))

        error = _save_logo(
            "notification_sound_file", settings, "notification_sound_file",
            filename_base="notification-custom",
            allowed_extensions=ALLOWED_SOUND_EXTENSIONS,
            invalid_format_message=_("Format nada harus MP3, WAV, OGG, atau M4A."),
        )
        if error:
            flash(error, "danger")
            return redirect(url_for("staff.admin_settings"))

        if request.form.get("remove_logo_square") == "1":
            _remove_logo(settings, "logo_square")

        if request.form.get("remove_logo_wide") == "1":
            _remove_logo(settings, "logo_wide")

        if request.form.get("remove_qris_image") == "1":
            _remove_logo(settings, "qris_image")

        if request.form.get("remove_notification_sound_file") == "1":
            _remove_logo(settings, "notification_sound_file")

        app_logo_choice = request.form.get("app_logo_choice", "square")
        login_logo_choice = request.form.get("login_logo_choice", "square")
        receipt_logo_choice = request.form.get("receipt_logo_choice", "wide")
        navbar_display = request.form.get("navbar_display", "both")

        settings.shop_name = shop_name
        settings.app_logo_choice = app_logo_choice if app_logo_choice in ("square", "wide") else "square"
        settings.login_logo_choice = login_logo_choice if login_logo_choice in ("square", "wide") else "square"
        settings.receipt_logo_choice = receipt_logo_choice if receipt_logo_choice in ("square", "wide") else "wide"
        settings.navbar_display = navbar_display if navbar_display in ("logo", "name", "both") else "both"

        settings.ppn_enabled = request.form.get("ppn_enabled") == "1"
        ppn_percentage = request.form.get("ppn_percentage", type=finite_float)
        if ppn_percentage is None:
            ppn_percentage = 0.0
        settings.ppn_percentage = max(0.0, min(100.0, ppn_percentage))

        paper_width = request.form.get("receipt_paper_width", "58")
        settings.receipt_paper_width = paper_width if paper_width in ("58", "80") else "58"
        settings.address = request.form.get("address", "").strip() or None
        settings.phone = request.form.get("phone", "").strip() or None
        settings.instagram = request.form.get("instagram", "").strip() or None
        settings.tiktok = request.form.get("tiktok", "").strip() or None
        settings.whatsapp = request.form.get("whatsapp", "").strip() or None
        settings.other_social = request.form.get("other_social", "").strip() or None
        settings.wifi_name = request.form.get("wifi_name", "").strip() or None
        settings.wifi_password = request.form.get("wifi_password", "").strip() or None

        settings.notification_enabled = request.form.get("notification_enabled") == "1"
        notification_sound = request.form.get("notification_sound", "church_bell")
        settings.notification_sound = (
            notification_sound if notification_sound in NOTIFICATION_SOUND_KEYS else "church_bell"
        )
        notification_volume = request.form.get("notification_volume", type=int)
        if notification_volume is None:
            notification_volume = 70
        settings.notification_volume = max(0, min(100, notification_volume))

        # Nada per role (Dapur/Kasir/Pelayan) - kosong berarti ikut nada
        # Owner di atas (lihat Settings.notification_sound_for_role()).
        for field in ("notification_sound_dapur", "notification_sound_kasir", "notification_sound_pelayan"):
            value = request.form.get(field, "").strip()
            setattr(settings, field, value if value in NOTIFICATION_SOUND_KEYS else None)

        # Ganti nama lantai (field floor_name_<id>) - ikut disimpan bareng
        # tombol "Simpan Pengaturan" yang sama, tidak perlu form terpisah.
        for floor in Floor.query.all():
            new_name = request.form.get(f"floor_name_{floor.id}", "").strip()
            if new_name:
                floor.name = new_name

        _sync_demo_backup(settings, values_before)

        db.session.commit()

        flash(_("Pengaturan toko berhasil disimpan."), "success")
        return redirect(url_for("staff.admin_settings"))

    floors = Floor.query.order_by(Floor.number).all()

    return render_template(
        "staff/settings_admin.html",
        settings=settings,
        floors=floors,
        demo_mode_unlocked=_dev_unlock_valid("demo_mode_unlocked"),
        factory_reset_unlocked=_dev_unlock_valid("factory_reset_unlocked"),
        **_system_tab_context(),
    )


@staff_bp.route("/admin/floors/add", methods=["POST"])
@roles_required(ROLE_OWNER)
def add_floor():
    """Tambah 1 lantai baru - nomornya otomatis lanjut dari yang paling
    besar. Nama dibiarkan kosong (bukan di-hardcode "Lantai N") supaya
    floor_display_name() tetap pakai fallback terjemahan sesuai bahasa
    aktif sampai Owner beneran mengisi nama custom lewat form utama -
    kalau di-hardcode di sini, nama defaultnya kebawa permanen dalam
    Bahasa Indonesia walau tampilan sedang bahasa Inggris.

    Tombol ini SATU form bareng field nama tiap lantai (lihat komentar
    struktur form di settings_admin.html) - kalau Owner sempat mengetik
    nama lantai lalu langsung klik "Tambah Lantai" (belum sempat klik
    "Simpan Pengaturan"), nama yang baru diketik itu ikut disimpan di
    sini dulu supaya tidak hilang begitu saja, baru lantai barunya
    ditambahkan."""

    for floor in Floor.query.all():
        new_name = request.form.get(f"floor_name_{floor.id}", "").strip()
        if new_name:
            floor.name = new_name

    # Ikut perhitungkan nomor lantai yang masih dipakai meja - kalau lantai
    # tertinggi sempat dihapus (mejanya tidak ikut terhapus), lantai baru
    # tidak boleh dapat nomor yang sama dan diam-diam "mewarisi" meja lama.
    last_number = max(
        db.session.query(db.func.max(Floor.number)).scalar() or 0,
        db.session.query(db.func.max(Table.floor)).scalar() or 0,
    )
    new_number = last_number + 1

    floor = Floor(number=new_number, name="")
    db.session.add(floor)
    db.session.commit()

    flash(_("Lantai %(number)s ditambahkan.", number=new_number), "success")
    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/floors/<int:floor_id>/delete", methods=["POST"])
@roles_required(ROLE_OWNER)
def delete_floor(floor_id):
    """Hapus 1 lantai dari daftar - meja yang kebetulan masih pakai nomor
    itu TIDAK ikut terhapus, cuma nama custom-nya hilang (otomatis balik
    ke label "Lantai N" polos) dan nomor itu tidak lagi jadi pilihan
    untuk meja baru."""

    floor = Floor.query.get_or_404(floor_id)
    display_name = floor_display_name(floor.number)
    db.session.delete(floor)
    db.session.commit()

    flash(_("%(name)s dihapus.", name=display_name), "success")
    return redirect(url_for("staff.admin_settings"))


# ============================================================
# SISTEM - BERSIHKAN CACHE & BACKUP DATABASE
# ============================================================

def _backup_folder():
    folder = current_app.config["BACKUP_FOLDER"]
    os.makedirs(folder, exist_ok=True)
    return folder


def _db_file_path():
    # Path absolut hasil resolusi Flask-SQLAlchemy (URI relatif sudah
    # dipetakan ke folder instance/), bukan potongan string URI mentah.
    return db.engine.url.database


def _folder_size_bytes(path):
    total = 0
    if not os.path.isdir(path):
        return total
    for entry in os.scandir(path):
        if entry.is_file():
            total += entry.stat().st_size
        elif entry.is_dir():
            total += _folder_size_bytes(entry.path)
    return total


def _cache_asset_breakdown():
    """Ukuran gambar yang di-cache browser (dibedakan per jenis) - dipakai
    supaya Owner tahu berapa MB yang sebenarnya "nyangkut" di device tamu/
    staf kalau belum di-refresh, bukan angka yang mengada-ada."""

    static_dir = os.path.join(current_app.root_path, "static")
    categories = [
        (_("Logo & Branding"), os.path.join(static_dir, "uploads", "branding")),
        (_("Foto Profil"), os.path.join(static_dir, "uploads", "avatars")),
        (_("Foto Menu"), os.path.join(static_dir, "uploads", "menu")),
        (_("QR Code Meja"), os.path.join(static_dir, "qrcodes")),
    ]

    sizes = [(label, _folder_size_bytes(path)) for label, path in categories]
    total_bytes = sum(size for _label, size in sizes)

    breakdown = [
        {
            "label": label,
            "mb": round(size / (1024 * 1024), 2),
            "percent": round((size / total_bytes) * 100) if total_bytes else 0,
        }
        for label, size in sizes
        if size > 0
    ]

    return {
        "total_mb": round(total_bytes / (1024 * 1024), 2),
        "categories": breakdown,
    }


@staff_bp.route("/admin/system")
@roles_required(ROLE_OWNER)
def admin_system():
    # Sistem digabung jadi tab di halaman Pengaturan - route lama ini
    # dipertahankan sebagai redirect saja, jaga-jaga ada bookmark/link lama.
    return redirect(url_for("staff.admin_settings"))


def _system_tab_context():
    """Data buat tab "Sistem" di halaman Pengaturan - cache, backup
    database, dan status Mode Demo. Dipisah jadi fungsi sendiri (bukan
    langsung di admin_settings()) supaya view function-nya tidak makin
    panjang campur aduk sama logika Pengaturan Toko yang lain."""

    backups = []
    for filename in os.listdir(_backup_folder()):
        if not filename.endswith(".zip"):
            continue
        path = os.path.join(_backup_folder(), filename)
        backups.append({
            "name": filename,
            "size_mb": round(os.path.getsize(path) / (1024 * 1024), 2),
            "created_at": datetime.fromtimestamp(os.path.getmtime(path)),
        })
    backups.sort(key=lambda b: b["created_at"], reverse=True)

    cache_info = _cache_asset_breakdown()
    demo_mode_active = Category.query.filter_by(is_demo=True).first() is not None
    maintenance_mode_active = get_settings().maintenance_mode

    if demo_mode_active:
        demo_stats = {
            "categories": Category.query.filter_by(is_demo=True).count(),
            "menu_items": MenuItem.query.filter_by(is_demo=True).count(),
            "ingredients": Ingredient.query.filter_by(is_demo=True).count(),
            "tables": Table.query.filter_by(is_demo=True).count(),
            "channels": OrderChannel.query.filter_by(is_demo=True).count(),
        }
    else:
        with open(_demo_fixture_path(), encoding="utf-8") as f:
            fixture = json.load(f)
        demo_stats = {
            "categories": len(fixture["categories"]),
            "menu_items": len(fixture["menu_items"]),
            "ingredients": len(fixture["ingredients"]),
            "tables": 10,
            "channels": len(DEMO_CHANNELS),
        }

    channels = OrderChannel.query.order_by(OrderChannel.sort_order, OrderChannel.id).all()

    return {
        "demo_stats": demo_stats,
        "backups": backups,
        "cache_info": cache_info,
        "demo_mode_active": demo_mode_active,
        "maintenance_mode_active": maintenance_mode_active,
        "channels": channels,
        "CHANNEL_PRICING_PERCENT": CHANNEL_PRICING_PERCENT,
        "CHANNEL_PRICING_MANUAL": CHANNEL_PRICING_MANUAL,
    }


# ============================================================
# PLATFORM OJOL (GrabFood, ShopeeFood, dst) - lihat OrderChannel di
# models.py. Dikelola sebagai daftar bebas (bukan field hardcode di
# Settings) supaya platform baru bisa ditambah kapan saja lewat
# Pengaturan > Platform Delivery tanpa perlu ubah kode sama sekali.
# ============================================================

@staff_bp.route("/admin/channels/add", methods=["POST"])
@roles_required(ROLE_OWNER)
def channel_add():
    name = request.form.get("name", "").strip()

    if not name:
        flash(_("Nama platform wajib diisi."), "danger")
        return redirect(url_for("staff.admin_settings"))

    pricing_mode = request.form.get("pricing_mode", CHANNEL_PRICING_PERCENT)
    if pricing_mode not in (CHANNEL_PRICING_PERCENT, CHANNEL_PRICING_MANUAL):
        pricing_mode = CHANNEL_PRICING_PERCENT

    markup_percent = min(request.form.get("markup_percent", type=finite_float) or 0, 1000.0)

    max_sort = db.session.query(db.func.max(OrderChannel.sort_order)).scalar() or 0
    channel = OrderChannel(
        name=name,
        pricing_mode=pricing_mode,
        markup_percent=max(0.0, markup_percent),
        sort_order=max_sort + 1,
    )
    db.session.add(channel)
    db.session.flush()

    error = _save_logo("logo", channel, "logo", filename_base=f"channel-logo-{channel.id}")
    if error:
        db.session.rollback()
        flash(error, "danger")
        return redirect(url_for("staff.admin_settings"))

    db.session.commit()
    flash(_("Platform %(name)s ditambahkan.", name=channel.name), "success")
    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/channels/<int:channel_id>/update", methods=["POST"])
@roles_required(ROLE_OWNER)
def channel_update(channel_id):
    channel = OrderChannel.query.get_or_404(channel_id)

    name = request.form.get("name", "").strip()
    if not name:
        flash(_("Nama platform wajib diisi."), "danger")
        return redirect(url_for("staff.admin_settings"))

    pricing_mode = request.form.get("pricing_mode", CHANNEL_PRICING_PERCENT)
    if pricing_mode not in (CHANNEL_PRICING_PERCENT, CHANNEL_PRICING_MANUAL):
        pricing_mode = CHANNEL_PRICING_PERCENT

    markup_percent = min(request.form.get("markup_percent", type=finite_float) or 0, 1000.0)

    error = _save_logo("logo", channel, "logo", filename_base=f"channel-logo-{channel.id}")
    if error:
        flash(error, "danger")
        return redirect(url_for("staff.admin_settings"))

    if request.form.get("remove_logo") == "1":
        _remove_logo(channel, "logo")

    channel.name = name
    channel.pricing_mode = pricing_mode
    channel.markup_percent = max(0.0, markup_percent)
    channel.is_active = request.form.get("is_active") == "1"

    db.session.commit()
    flash(_("Platform %(name)s diperbarui.", name=channel.name), "success")
    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/channels/<int:channel_id>/delete", methods=["POST"])
@roles_required(ROLE_OWNER)
def channel_delete(channel_id):
    channel = OrderChannel.query.get_or_404(channel_id)

    # Pesanan lama yang pernah pakai platform ini TIDAK ikut dihapus -
    # cukup lepas kaitannya (channel_id jadi kosong) supaya riwayat &
    # laporan lama tetap ada, cuma label platformnya lewat display_label
    # fallback ke tulisan umum "Ojek Online".
    for order in Order.query.filter_by(channel_id=channel.id).all():
        order.channel_id = None

    _remove_logo(channel, "logo")
    db.session.delete(channel)
    db.session.commit()

    flash(_("Platform %(name)s dihapus.", name=channel.name), "success")
    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/channels/<int:channel_id>/prices", methods=["GET", "POST"])
@roles_required(ROLE_OWNER)
def channel_prices(channel_id):
    channel = OrderChannel.query.get_or_404(channel_id)

    if request.method == "POST":
        valid_menu_ids = {item_id for (item_id,) in db.session.query(MenuItem.id).all()}

        for key, value in request.form.items():
            if not key.startswith("price_"):
                continue

            raw_id = key[len("price_"):]
            if not raw_id.isdigit() or int(raw_id) not in valid_menu_ids:
                continue

            menu_item_id = int(raw_id)
            price_text = value.strip()

            row = MenuItemChannelPrice.query.filter_by(
                channel_id=channel.id, menu_item_id=menu_item_id
            ).first()

            if not price_text:
                # Dikosongkan = balik pakai harga toko biasa (fallback),
                # jadi baris override-nya dihapus saja.
                if row:
                    db.session.delete(row)
                continue

            try:
                price = digits_int(price_text)
            except ValueError:
                continue
            if row:
                row.price = price
            else:
                db.session.add(
                    MenuItemChannelPrice(
                        channel_id=channel.id, menu_item_id=menu_item_id, price=price
                    )
                )

        db.session.commit()
        flash(_("Harga khusus %(name)s disimpan.", name=channel.name), "success")
        return redirect(url_for("staff.channel_prices", channel_id=channel.id))

    categories = Category.query.order_by(Category.order, Category.name).all()
    custom_prices = {
        row.menu_item_id: row.price
        for row in MenuItemChannelPrice.query.filter_by(channel_id=channel.id).all()
    }

    return render_template(
        "staff/channel_prices.html",
        channel=channel,
        categories=categories,
        custom_prices=custom_prices,
    )


def _demo_fixture_path():
    return os.path.join(current_app.root_path, "demo_data", "fixture.json")


def _demo_settings_fixture_path():
    return os.path.join(current_app.root_path, "demo_data", "settings_fixture.json")


SETTINGS_DEMO_TEXT_FIELDS = (
    "shop_name",
    "address",
    "phone",
    "instagram",
    "tiktok",
    "whatsapp",
    "other_social",
    "wifi_name",
    "wifi_password",
    "notification_sound_dapur",
    "notification_sound_kasir",
    "notification_sound_pelayan",
)


def _apply_demo_settings():
    """Timpa identitas toko (nama, alamat, kontak, logo) di Settings
    dengan isi contoh dari app/demo_data/settings_fixture.json - nilai
    ASLI-nya disimpan dulu ke Settings.demo_settings_backup (JSON)
    supaya bisa dikembalikan persis lewat _restore_real_settings()
    begitu Mode Demo dimatikan. Logo diganti dengan MENYALIN file demo
    ke folder upload (bukan menghapus logo asli), jadi logo asli tetap
    utuh di disk sepanjang demo berlangsung."""

    settings = get_settings()

    with open(_demo_settings_fixture_path(), encoding="utf-8") as f:
        fixture = json.load(f)

    backup = {field: getattr(settings, field) for field in SETTINGS_DEMO_TEXT_FIELDS}
    backup["logo_square"] = settings.logo_square
    backup["logo_wide"] = settings.logo_wide
    settings.demo_settings_backup = json.dumps(backup)

    for field in SETTINGS_DEMO_TEXT_FIELDS:
        setattr(settings, field, fixture.get(field))

    branding_dir = os.path.join(current_app.static_folder, "uploads", "branding")
    os.makedirs(branding_dir, exist_ok=True)
    demo_branding_dir = os.path.join(current_app.root_path, "demo_data", "branding")

    for field, demo_filename in fixture.get("logo_files", {}).items():
        src = os.path.join(demo_branding_dir, demo_filename)
        if os.path.exists(src):
            dest_name = f"demo-{demo_filename}"
            shutil.copyfile(src, os.path.join(branding_dir, dest_name))
            setattr(settings, field, dest_name)


DEMO_BACKUP_FIELDS = SETTINGS_DEMO_TEXT_FIELDS + ("logo_square", "logo_wide")


def _sync_demo_backup(settings, values_before):
    """Selama Mode Demo aktif, identitas toko yang tampil adalah versi
    CONTOH, sedangkan versi asli disimpan di demo_settings_backup dan
    dipulihkan saat demo dimatikan. Kalau Owner mengubah identitas toko
    di tengah demo, perubahan itu dianggap untuk toko ASLI - ikut
    ditulis ke backup supaya tidak tertimpa nilai lama saat demo
    dimatikan. Field yang tidak diubah tetap memakai nilai asli di
    backup."""

    if not settings.demo_settings_backup:
        return

    try:
        backup = json.loads(settings.demo_settings_backup)
    except ValueError:
        return

    for field in SETTINGS_DEMO_TEXT_FIELDS:
        new_value = getattr(settings, field)
        if new_value != values_before.get(field):
            backup[field] = new_value

    for field in ("logo_square", "logo_wide"):
        old_value = values_before.get(field)
        new_value = getattr(settings, field)
        if new_value == old_value:
            continue
        if new_value is None and (old_value or "").startswith("demo-"):
            # Cuma logo CONTOH yang dihapus - logo asli di backup aman.
            continue
        real_old = backup.get(field)
        if real_old and real_old != new_value:
            _delete_file_after_commit(os.path.join(_branding_upload_folder(), real_old))
        backup[field] = new_value

    settings.demo_settings_backup = json.dumps(backup)


def _restore_real_settings():
    """Lawan dari _apply_demo_settings() - kembalikan nama toko, kontak,
    dan logo ke nilai asli dari sebelum Mode Demo aktif. File logo demo
    (prefix "demo-") ikut dihapus dari folder upload, file logo ASLI
    tidak pernah tersentuh sama sekali sejak awal."""

    settings = get_settings()
    if not settings.demo_settings_backup:
        return

    backup = json.loads(settings.demo_settings_backup)

    branding_dir = os.path.join(current_app.static_folder, "uploads", "branding")
    for field in ("logo_square", "logo_wide"):
        current_filename = getattr(settings, field)
        if current_filename and current_filename.startswith("demo-"):
            _delete_file_after_commit(os.path.join(branding_dir, current_filename))

    for field, value in backup.items():
        setattr(settings, field, value)

    settings.demo_settings_backup = None


def _random_datetime_between(start, end):
    delta = end - start
    seconds = random.uniform(0, delta.total_seconds())
    return start + timedelta(seconds=seconds)


def _demo_combo_items(drinks, foods, target_min=100_000, target_max=300_000, tries=30):
    """Pilih 2-3 item (minuman + makanan, kadang 1 tambahan) acak sampai
    subtotalnya masuk rentang target - dipakai buat contoh "1 meja pesan
    paket lengkap". Kalau 30x percobaan tidak ada yang pas, baris
    terakhir yang dicoba tetap dipakai apa adanya (tidak penting-penting
    amat presisi buat data contoh)."""

    picks = []
    for _ in range(tries):
        picks = [
            (random.choice(drinks), random.randint(1, 2)),
            (random.choice(foods), random.randint(1, 2)),
        ]
        if random.random() < 0.4:
            picks.append((random.choice(drinks + foods), 1))
        subtotal = sum(item.price * qty for item, qty in picks)
        if target_min <= subtotal <= target_max:
            break
    return picks


def _demo_make_order(
    order_type, items, table=None, channel=None, status="pending",
    is_paid=False, paid_at=None, deduct_stock=True,
):
    """Helper bikin 1 order contoh (dipakai berulang oleh
    _seed_demo_orders()) - urus item/snapshot harga/opsional potong
    stok/opsional tandai lunas dalam satu tempat, supaya tiap skenario
    demo (meja belum bayar, transaksi lunas, dst.) tidak menulis ulang
    logika yang sama."""

    order = Order(
        table_id=table.id if table else None,
        order_type=order_type,
        channel_id=channel.id if channel else None,
        source="staff",
        status="served" if is_paid else status,
        pin=generate_order_pin(),
    )

    total = 0
    for menu_item, qty in items:
        price = channel.price_for(menu_item) if channel else menu_item.price
        order.items.append(OrderItem(
            menu_item_id=menu_item.id,
            name_snapshot=menu_item.name,
            price_snapshot=price,
            quantity=qty,
        ))
        total += price * qty

    # PENTING: semua field lunas (is_paid/status/paid_at/dst.) HARUS
    # sudah di-set SEBELUM db.session.add()/check_and_deduct_stock() -
    # check_and_deduct_stock() memanggil query yang memicu autoflush
    # SQLAlchemy, jadi kalau field-field ini di-set BELAKANGAN, order
    # keburu tersimpan ke DB dalam keadaan belum-lunas (bisa tabrakan
    # sama kunci "1 order belum lunas per meja" kalau meja itu sudah
    # dipakai order lain yang sengaja dibiarkan belum bayar).
    if is_paid:
        order.is_paid = True
        order.payment_method = random.choice(PAYMENT_METHODS)
        order.served_by = "demo"
        order.paid_at = paid_at or datetime.now()
        order.created_at = order.paid_at
        if order.payment_method == "cash":
            order.cash_received = total
            order.change_amount = 0

    db.session.add(order)

    if deduct_stock:
        # Stok bahan baku demo sengaja dibuat cukup besar (lihat fixture.json
        # & 2 bahan kemasan di _seed_demo_data()) supaya beberapa lusin
        # pesanan contoh ini tidak sampai kehabisan stok di tengah proses
        # seed - kalaupun ada yang gagal (mis. race jarak-jauh yang tidak
        # relevan di sini), diabaikan saja, pesanan tetap dibuat tanpa
        # potongan stok daripada seluruh Mode Demo gagal aktif gara-gara ini.
        check_and_deduct_stock(order)

    return order, total


def _seed_demo_filler_orders(start_dt, end_dt, target_total, demo_items, demo_tables, demo_channels):
    """Isi riwayat transaksi lunas acak dalam 1 rentang waktu sampai
    totalnya kira-kira mencapai target_total - dipakai buat contoh
    angka Laporan Penjualan (bulanan/tahunan) yang masuk akal besarnya,
    bukan cuma beberapa transaksi kosong. Sengaja TIDAK lewat
    check_and_deduct_stock (bisa ratusan baris, dan tujuannya cuma
    angka laporan, bukan demo potong stok - itu sudah diwakili pesanan
    lain di _seed_demo_orders())."""

    running_total = 0
    count = 0
    max_orders = 10_000  # jaring pengaman - jangan sampai muter tanpa henti

    while running_total < target_total and count < max_orders:
        n_items = random.choices([1, 2, 3], weights=[5, 3, 2])[0]
        picks = [(random.choice(demo_items), random.randint(1, 4)) for _ in range(n_items)]

        order_type = random.choices(
            [ORDER_TYPE_DINE_IN, ORDER_TYPE_TAKEAWAY, ORDER_TYPE_OJOL],
            weights=[55, 35, 10],
        )[0]
        table = random.choice(demo_tables) if order_type == ORDER_TYPE_DINE_IN else None
        channel = random.choice(demo_channels) if order_type == ORDER_TYPE_OJOL else None

        _, total = _demo_make_order(
            order_type, picks, table=table, channel=channel,
            is_paid=True, paid_at=_random_datetime_between(start_dt, end_dt),
            deduct_stock=False,
        )
        running_total += total
        count += 1

    return running_total, count


def _seed_demo_orders():
    """Isi contoh riwayat pesanan buat Mode Demo - dipanggil di akhir
    _seed_demo_data() setelah menu/meja/channel/bahan kemasan demo
    selesai dibuat. Tidak ada kolom is_demo di Order/OrderItem - tetap
    bisa dibersihkan otomatis oleh _clear_demo_data() karena SEMUA order
    di sini memakai menu_item_id yang menunjuk ke MenuItem is_demo=True,
    dan pembersihan yang sudah ada menghapus tiap order yang mengandung
    item demo (lihat _clear_demo_data(), bagian demo_item_ids)."""

    demo_tables = Table.query.filter_by(is_demo=True).order_by(Table.label).all()
    demo_channels = OrderChannel.query.filter_by(is_demo=True).all()
    demo_items = MenuItem.query.filter_by(is_demo=True).all()
    drinks = [m for m in demo_items if m.category.name in DEMO_DRINK_CATEGORIES]
    foods = [m for m in demo_items if m.category.name in DEMO_FOOD_CATEGORIES]

    packaged_items = (
        MenuItem.query.join(MenuItemIngredient, MenuItemIngredient.menu_item_id == MenuItem.id)
        .join(Ingredient, MenuItemIngredient.ingredient_id == Ingredient.id)
        .filter(MenuItem.is_demo.is_(True), Ingredient.name.in_(
            [DEMO_PACKAGING_CUP_NAME, DEMO_PACKAGING_BOX_NAME]
        ))
        .distinct()
        .all()
    ) or demo_items[:6]

    today = datetime.now()
    today_start = datetime.combine(today.date(), time.min)

    # 1) 5 meja "Makan di Tempat" BELUM BAYAR - masih menunggu di kasir.
    #    3 di antaranya statusnya "ready" (dapur sudah selesai, tinggal
    #    diantar pelayan), 2 sisanya masih "pending" (baru diterima).
    unpaid_tables = demo_tables[:5]
    for i, table in enumerate(unpaid_tables):
        items = _demo_combo_items(drinks, foods)
        status = "ready" if i < 3 else "pending"
        _demo_make_order(ORDER_TYPE_DINE_IN, items, table=table, status=status)

    # 2) 5 pesanan "Bawa Pulang" LUNAS hari ini - paket minum+makan.
    for _ in range(5):
        items = _demo_combo_items(drinks, foods)
        _demo_make_order(ORDER_TYPE_TAKEAWAY, items, is_paid=True, paid_at=_random_datetime_between(today_start, today))

    # 3) 5 pesanan contoh yang menu-nya memakai bahan kemasan (gelas cup/
    #    kotak styrofoam) - diacak jenis pesanannya (dine-in/bawa pulang/
    #    ojol) supaya potongan stok kemasan kelihatan kepakai lintas
    #    platform, bukan cuma 1 jenis pesanan saja.
    for _ in range(5):
        item = random.choice(packaged_items)
        order_type = random.choice([ORDER_TYPE_DINE_IN, ORDER_TYPE_TAKEAWAY, ORDER_TYPE_OJOL])
        table = random.choice(demo_tables) if order_type == ORDER_TYPE_DINE_IN else None
        channel = random.choice(demo_channels) if order_type == ORDER_TYPE_OJOL else None
        _demo_make_order(
            order_type, [(item, random.randint(1, 2))], table=table, channel=channel,
            is_paid=True, paid_at=_random_datetime_between(today_start, today),
        )

    # 4) Tambahan pesanan lunas hari ini (jenis & isi acak) supaya
    #    "Riwayat Hari Ini" di Kasir menampilkan 20 transaksi siap
    #    cetak ulang - bukan cuma 10 dari langkah 2 & 3 di atas.
    for _ in range(10):
        n_items = random.randint(1, 2)
        items = [(random.choice(demo_items), random.randint(1, 2)) for _ in range(n_items)]
        order_type = random.choices(
            [ORDER_TYPE_DINE_IN, ORDER_TYPE_TAKEAWAY, ORDER_TYPE_OJOL], weights=[50, 35, 15],
        )[0]
        table = random.choice(demo_tables) if order_type == ORDER_TYPE_DINE_IN else None
        channel = random.choice(demo_channels) if order_type == ORDER_TYPE_OJOL else None
        _demo_make_order(
            order_type, items, table=table, channel=channel,
            is_paid=True, paid_at=_random_datetime_between(today_start, today),
        )

    db.session.flush()
    today_revenue = (
        db.session.query(func.coalesce(func.sum(OrderItem.price_snapshot * OrderItem.quantity), 0))
        .join(Order, OrderItem.order_id == Order.id)
        .filter(Order.is_paid.is_(True), Order.paid_at >= today_start)
        .scalar()
    )

    # 5) Sisa hari di bulan berjalan (sebelum hari ini) diisi transaksi
    #    acak sampai total bulan ini kira-kira Rp30 juta, dan 8 bulan
    #    sebelumnya (Januari s/d bulan lalu) sampai total tahun ini
    #    kira-kira Rp400 juta - buat contoh angka Laporan Penjualan yang
    #    masuk akal besarnya. Random murni, tidak perlu presisi.
    month_start = today.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if month_start < today_start:
        _seed_demo_filler_orders(
            month_start, today_start, max(0, 30_000_000 - today_revenue),
            demo_items, demo_tables, demo_channels,
        )

    year_start = today.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
    if year_start < month_start:
        _seed_demo_filler_orders(
            year_start, month_start, 370_000_000,
            demo_items, demo_tables, demo_channels,
        )


def _seed_demo_data():
    """Isi menu + inventory dummy dari app/demo_data/fixture.json (foto
    ikut disalin dari app/demo_data/photos/) - ditandai is_demo=True
    supaya nanti bisa dihapus lagi tanpa menyentuh data asli toko."""

    with open(_demo_fixture_path(), encoding="utf-8") as f:
        fixture = json.load(f)

    categories_by_name = {}
    for c in fixture["categories"]:
        cat = Category(name=c["name"], order=c["order"], is_demo=True)
        db.session.add(cat)
        categories_by_name[c["name"]] = cat

    ingredients_by_name = {}
    for i in fixture["ingredients"]:
        ing = Ingredient(
            name=i["name"],
            unit=i["unit"],
            stock_quantity=i["stock_quantity"],
            low_stock_threshold=i["low_stock_threshold"],
            is_demo=True,
        )
        db.session.add(ing)
        ingredients_by_name[i["name"]] = ing

    db.session.flush()

    upload_dir = os.path.join(current_app.static_folder, "uploads", "menu")
    os.makedirs(upload_dir, exist_ok=True)
    demo_photos_dir = os.path.join(current_app.root_path, "demo_data", "photos")

    for m in fixture["menu_items"]:
        item = MenuItem(
            category=categories_by_name[m["category"]],
            name=m["name"],
            price=m["price"],
            is_available=True,
            is_demo=True,
        )
        db.session.add(item)
        db.session.flush()

        if m["photo"]:
            src = os.path.join(demo_photos_dir, m["photo"])
            if os.path.exists(src):
                ext = os.path.splitext(m["photo"])[1]
                dest_name = f"menu-{item.id}{ext}"
                shutil.copyfile(src, os.path.join(upload_dir, dest_name))
                item.photo = dest_name

        for row in m["recipe"]:
            db.session.add(
                MenuItemIngredient(
                    menu_item_id=item.id,
                    ingredient_id=ingredients_by_name[row["ingredient"]].id,
                    quantity_used=row["qty"],
                )
            )

    # 10 meja demo di Lantai 1 - kode pakai prefix "demo-" (bukan
    # "postab-" seperti meja asli) supaya DIJAMIN tidak pernah tabrakan
    # sama kode meja toko yang sungguhan, walau tokonya sudah punya meja
    # 1-10 sendiri. Owner boleh coba alur pesanan sungguhan di meja ini
    # buat demo - order yang kebentuk ikut dibersihkan di _clear_demo_data().
    #
    # QR code-nya digenerate di sini juga (sama seperti meja asli lewat
    # table_map()) - sebelumnya baris Table cuma ditambahkan ke DB tanpa
    # pernah memanggil _generate_table_qr(), jadi file gambar QR-nya
    # tidak pernah ada dan modal "Detail Meja" di Mode Demo selalu
    # gagal muat gambar (terlihat rusak/crash).
    for n in range(1, 11):
        table = Table(code=f"demo-{n:02d}", label=f"Meja {n}", floor=1, is_demo=True)
        db.session.add(table)
        _generate_table_qr(table)

    # 1 lantai contoh (Lantai 1) - cuma dibuat kalau toko belum punya
    # Floor bernomor 1 sendiri, supaya tidak bentrok sama data asli
    # (Floor.number unique). 10 meja demo di atas semuanya sudah pakai
    # floor=1, jadi baris ini cuma nempelkan NAMA-nya saja.
    if not Floor.query.filter_by(number=1).first():
        db.session.add(Floor(number=1, name=DEMO_FLOOR_NAME, is_demo=True))

    # 2 bahan baku kemasan (gelas & kotak sekali pakai) - ditempel ke
    # beberapa menu contoh (bukan bagian resep asli fixture.json) supaya
    # potongan stok kemasan kelihatan nyata dipakai, bukan cuma bahan
    # baku "isi" minuman/makanannya saja.
    packaging_cup = Ingredient(
        name=DEMO_PACKAGING_CUP_NAME, unit="pcs", stock_quantity=300, low_stock_threshold=50, is_demo=True,
    )
    packaging_box = Ingredient(
        name=DEMO_PACKAGING_BOX_NAME, unit="pcs", stock_quantity=300, low_stock_threshold=50, is_demo=True,
    )
    db.session.add_all([packaging_cup, packaging_box])
    db.session.flush()

    demo_drink_items = [m for m in categories_by_name["Kopi"].items if m.is_demo][:3] if "Kopi" in categories_by_name else []
    demo_food_items = [m for m in categories_by_name["Makanan Berat"].items if m.is_demo][:3] if "Makanan Berat" in categories_by_name else []
    for item in demo_drink_items:
        db.session.add(MenuItemIngredient(menu_item_id=item.id, ingredient_id=packaging_cup.id, quantity_used=1))
    for item in demo_food_items:
        db.session.add(MenuItemIngredient(menu_item_id=item.id, ingredient_id=packaging_box.id, quantity_used=1))
    db.session.flush()

    # 3 platform delivery demo - markup persentase mengikuti kisaran komisi
    # riil yang berlaku ke mitra reguler (per September 2026, bukan
    # merchant preferred/strategis yang komisinya bisa lebih rendah):
    # GoFood ~20%, GrabFood ~30%, ShopeeFood ~20%. Dipakai Owner buat
    # contoh sebelum ganti sendiri ke platform & angka yang sungguhan
    # dipakai tokonya. Logo lencana contoh ikut disalin dari
    # app/demo_data/branding/ (sama pola-nya dengan _apply_demo_settings()).
    branding_dir = os.path.join(current_app.static_folder, "uploads", "branding")
    os.makedirs(branding_dir, exist_ok=True)
    demo_branding_dir = os.path.join(current_app.root_path, "demo_data", "branding")

    for i, (name, markup, badge_filename) in enumerate(DEMO_CHANNELS, start=1):
        channel = OrderChannel(
            name=name,
            pricing_mode=CHANNEL_PRICING_PERCENT,
            markup_percent=markup,
            is_active=True,
            sort_order=i,
            is_demo=True,
        )
        db.session.add(channel)
        db.session.flush()

        src = os.path.join(demo_branding_dir, badge_filename)
        if os.path.exists(src):
            dest_name = f"channel-logo-{channel.id}.png"
            shutil.copyfile(src, os.path.join(branding_dir, dest_name))
            channel.logo = dest_name

    _apply_demo_settings()
    _seed_demo_orders()

    db.session.commit()


def _demo_clear_blockers():
    """Data ASLI (is_demo=False) yang masih nyangkut ke data demo - menu
    asli yang ditaruh di kategori demo, atau resep menu asli yang pakai
    bahan baku demo. Kalau ada, Mode Demo tidak boleh dimatikan dulu:
    menghapus kategori/bahan demo akan gagal (NOT NULL) atau merusak
    data asli itu. Balikin list nama yang perlu dipindah/diubah Owner."""

    blockers = []

    for item in (
        MenuItem.query.join(Category, MenuItem.category_id == Category.id)
        .filter(Category.is_demo.is_(True), MenuItem.is_demo.is_(False))
        .all()
    ):
        blockers.append(item.name)

    for row in (
        MenuItemIngredient.query
        .join(Ingredient, MenuItemIngredient.ingredient_id == Ingredient.id)
        .join(MenuItem, MenuItemIngredient.menu_item_id == MenuItem.id)
        .filter(Ingredient.is_demo.is_(True), MenuItem.is_demo.is_(False))
        .all()
    ):
        blockers.append(row.menu_item.name)

    return sorted(set(blockers))


def _clear_demo_data():
    """Hapus semua baris is_demo=True (menu, kategori, bahan baku, meja)
    beserta file fotonya, dan kembalikan identitas toko (Settings) ke
    nilai asli lewat _restore_real_settings() - data asli toko
    (is_demo=False) sama sekali tidak disentuh.

    Dihapus satu-satu lewat db.session.delete() (bukan bulk .delete())
    supaya cascade "all, delete-orphan" di MenuItem.recipe & Order.items
    beneran jalan dan ikut membersihkan baris turunannya - bulk delete
    lewat query itu bypass ORM jadi tidak memicu cascade sama sekali.

    Meja demo (Table.is_demo=True) dihapus PALING TERAKHIR, dan order apa
    pun yang sempat dibuat Owner di meja itu saat demo (buat coba alur
    pesanan) dihapus duluan - meja tidak bisa dihapus selagi masih ada
    order yang mengacu ke situ (foreign key)."""

    upload_dir = os.path.join(current_app.static_folder, "uploads", "menu")
    demo_items = MenuItem.query.filter_by(is_demo=True).all()
    demo_item_ids = [item.id for item in demo_items]

    # Pesanan yang berisi menu demo (termasuk Bawa Pulang yang tidak punya
    # meja/platform demo) ikut dihapus - kalau tidak, penjualan contoh ini
    # nyangkut selamanya di Laporan penjualan asli.
    if demo_item_ids:
        demo_order_ids = select(OrderItem.order_id).where(OrderItem.menu_item_id.in_(demo_item_ids))
        for order in Order.query.filter(Order.id.in_(demo_order_ids)).all():
            db.session.delete(order)

    for item in demo_items:
        if item.photo:
            _delete_file_after_commit(os.path.join(upload_dir, item.photo))
        db.session.delete(item)

    for cat in Category.query.filter_by(is_demo=True).all():
        db.session.delete(cat)

    for ing in Ingredient.query.filter_by(is_demo=True).all():
        db.session.delete(ing)

    demo_tables = Table.query.filter_by(is_demo=True).all()
    demo_table_ids = [t.id for t in demo_tables]
    for order in Order.query.filter(Order.table_id.in_(demo_table_ids)).all():
        db.session.delete(order)

    qr_dir = os.path.join(current_app.root_path, "static", "qrcodes")
    for table in demo_tables:
        _delete_file_after_commit(os.path.join(qr_dir, f"{table.code}.png"))
        db.session.delete(table)

    # Platform delivery demo (GoFood/GrabFood/ShopeeFood contoh) - order
    # yang sempat dibuat lewat platform ini TIDAK punya table_id (lihat
    # ORDER_TYPE_OJOL di new_order()), jadi tidak ikut kena pembersihan
    # order berbasis meja di atas - harus dicari lewat channel_id sendiri.
    demo_channels = OrderChannel.query.filter_by(is_demo=True).all()
    demo_channel_ids = [c.id for c in demo_channels]
    for order in Order.query.filter(Order.channel_id.in_(demo_channel_ids)).all():
        db.session.delete(order)
    for channel in demo_channels:
        _remove_logo(channel, "logo")
        db.session.delete(channel)

    for floor in Floor.query.filter_by(is_demo=True).all():
        db.session.delete(floor)

    _restore_real_settings()

    db.session.commit()


def _check_dev_pin():
    """Cek PIN developer dengan batas percobaan salah per user - tanpa ini
    Owner (yang justru mau dicegah) bisa nebak PIN pakai script."""

    rate_key = f"dev-pin:{current_user.id}"
    if is_blocked(rate_key, DEV_PIN_RATE_LIMIT, DEV_PIN_RATE_WINDOW_SECONDS):
        return False

    if check_password_hash(DEV_PIN_HASH, request.form.get("pin", "")):
        return True

    record_failure(rate_key, DEV_PIN_RATE_WINDOW_SECONDS)
    return False


def _dev_unlock_valid(session_key):
    """True kalau PIN developer untuk fitur ini dibuka kurang dari
    DEV_UNLOCK_TTL_SECONDS yang lalu. Nilai lama (True, dari sebelum ada
    batas waktu) dianggap sudah kedaluwarsa."""

    unlocked_at = session.get(session_key)
    if isinstance(unlocked_at, bool) or not isinstance(unlocked_at, (int, float)):
        return False
    return datetime.now().timestamp() - unlocked_at < DEV_UNLOCK_TTL_SECONDS


@staff_bp.route("/admin/system/demo-mode/unlock", methods=["POST"])
@roles_required(ROLE_OWNER)
def demo_mode_unlock():
    """Verifikasi PIN Mode Demo - kalau cocok, tandai "terbuka" di session
    (server-side, bukan sekadar tampilan) supaya request langsung ke
    demo_mode_toggle() tanpa lewat tombol ini pun tetap kena tolak."""

    if _check_dev_pin():
        session["demo_mode_unlocked"] = datetime.now().timestamp()
        return {"ok": True}

    return {"ok": False}, 403


@staff_bp.route("/admin/system/demo-mode/lock", methods=["POST"])
@roles_required(ROLE_OWNER)
def demo_mode_lock():
    """Kunci lagi Mode Demo TANPA logout - lawan dari demo_mode_unlock().
    Tidak butuh PIN buat mengunci (cuma buat buka), tapi tetap harus
    login sebagai Owner supaya orang random tidak bisa iseng ngunci
    punya orang lain."""

    session.pop("demo_mode_unlocked", None)
    return {"ok": True}


@staff_bp.route("/admin/system/demo-mode/toggle", methods=["POST"])
@roles_required(ROLE_OWNER)
def demo_mode_toggle():
    if not _dev_unlock_valid("demo_mode_unlocked"):
        flash(_("Kunci PIN developer sudah kedaluwarsa. Buka kunci lagi dengan PIN."), "warning")
        return redirect(url_for("staff.admin_settings"))

    is_active = Category.query.filter_by(is_demo=True).first() is not None

    if is_active:
        blockers = _demo_clear_blockers()
        if blockers:
            flash(
                _(
                    "Mode Demo belum bisa dimatikan: menu asli berikut masih memakai kategori/bahan baku demo - pindahkan kategori atau ubah resepnya dulu: %(names)s.",
                    names=", ".join(blockers),
                ),
                "danger",
            )
            return redirect(url_for("staff.admin_settings"))
        _clear_demo_data()
        flash(_("Mode Demo dimatikan - semua data dummy sudah dihapus."), "success")
    else:
        _seed_demo_data()
        flash(_("Mode Demo diaktifkan - menu & inventory dummy sudah dimuat."), "success")

    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/system/maintenance-mode/toggle", methods=["POST"])
@roles_required(ROLE_OWNER)
def maintenance_mode_toggle():
    settings = get_settings()
    settings.maintenance_mode = not settings.maintenance_mode
    db.session.commit()

    if settings.maintenance_mode:
        flash(
            _("Mode Perbaikan diaktifkan - hanya akun Owner yang masih bisa mengakses aplikasi."),
            "success",
        )
    else:
        flash(_("Mode Perbaikan dimatikan - aplikasi kembali normal untuk semua pengguna."), "success")

    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/system/factory-reset/unlock", methods=["POST"])
@roles_required(ROLE_OWNER)
def factory_reset_unlock():
    """Sama persis konsepnya dengan demo_mode_unlock() - PIN dicek di
    server, status "terbuka" disimpan di session sendiri (TERPISAH dari
    demo_mode_unlocked) supaya buka kunci Mode Demo tidak otomatis buka
    kunci Reset Pabrik juga, walau PIN-nya sama."""

    if _check_dev_pin():
        session["factory_reset_unlocked"] = datetime.now().timestamp()
        return {"ok": True}

    return {"ok": False}, 403


@staff_bp.route("/admin/system/factory-reset/lock", methods=["POST"])
@roles_required(ROLE_OWNER)
def factory_reset_lock():
    session.pop("factory_reset_unlocked", None)
    return {"ok": True}


@staff_bp.route("/admin/system/factory-reset/execute", methods=["POST"])
@roles_required(ROLE_OWNER)
def factory_reset_execute():
    if not _dev_unlock_valid("factory_reset_unlocked"):
        flash(_("Kunci PIN developer sudah kedaluwarsa. Buka kunci lagi dengan PIN."), "warning")
        return redirect(url_for("staff.admin_settings"))

    _factory_reset()
    session.pop("factory_reset_unlocked", None)

    flash(
        _(
            "Reset ke mode pabrik selesai - semua data toko sudah dikosongkan, "
            "kecuali daftar pengguna dan pengaturan Notifikasi."
        ),
        "success",
    )
    return redirect(url_for("staff.admin_settings"))


def _factory_reset():
    """Kosongkan SEMUA data operasional toko (kategori, menu, resep,
    bahan baku, meja, lantai, riwayat pesanan/login) dan kembalikan
    Pengaturan Toko ke kondisi kosong seperti baru install - KECUALI
    tabel users (username, role, izin staf) DAN pengaturan Notifikasi
    (aktif/tidak, suara, volume per peran) - keduanya sengaja tidak
    disentuh sama sekali, sama seperti users, supaya konfigurasi yang
    sudah diatur pemilik tidak ikut hilang tiap kali reset. Urutan
    hapus mengikuti dependensi FK (baris anak dulu baru induk) lewat
    bulk delete, bukan lewat relationship cascade ORM, supaya aman
    terlepas dari cascade aktif atau tidak di DB yang dipakai."""

    # Platform delivery (OrderChannel) & harga khususnya - SENGAJA ikut
    # dihapus juga (sebelumnya kelewat), soalnya kalau tidak, channel
    # demo (GoFood/GrabFood/ShopeeFood beserta markup-nya) bisa nyangkut
    # hidup terus setelah reset dan diam-diam kepakai di pesanan asli.
    for channel in OrderChannel.query.all():
        _remove_logo(channel, "logo")
    MenuItemChannelPrice.query.delete()
    OrderStockDeduction.query.delete()
    CancelledOrderLog.query.delete()
    OrderItem.query.delete()
    Order.query.delete()
    OrderChannel.query.delete()
    MenuItemIngredient.query.delete()
    MenuItem.query.delete()
    Category.query.delete()
    Ingredient.query.delete()
    Table.query.delete()
    Floor.query.delete()
    LoginLog.query.delete()

    # Foto menu yang filenya sudah yatim (baris MenuItem-nya baru dihapus
    # di atas) - dibersihkan dari disk juga, bukan cuma dari DB.
    # QR code meja juga - semua meja sudah dihapus di atas.
    menu_upload_folder = _upload_folder("menu")
    qrcodes_folder = os.path.join(current_app.static_folder, "qrcodes")
    for folder in (menu_upload_folder, qrcodes_folder):
        if not os.path.isdir(folder):
            continue
        for filename in os.listdir(folder):
            file_path = os.path.join(folder, filename)
            if os.path.isfile(file_path):
                _delete_file_after_commit(file_path)

    settings = get_settings()
    # Kalau reset dilakukan saat Mode Demo aktif, logo ASLI dari sebelum
    # demo (disimpan di demo_settings_backup) juga ikut dibersihkan.
    if settings.demo_settings_backup:
        try:
            backup = json.loads(settings.demo_settings_backup)
        except ValueError:
            backup = {}
        for field in ("logo_square", "logo_wide", "qris_image"):
            filename = backup.get(field)
            if filename:
                _delete_file_after_commit(os.path.join(_branding_upload_folder(), filename))
    _remove_logo(settings, "logo_square")
    _remove_logo(settings, "logo_wide")
    _remove_logo(settings, "qris_image")
    # notification_sound_file SENGAJA tidak dihapus - lihat catatan di
    # docstring, pengaturan Notifikasi (termasuk file suara custom-nya
    # kalau ada) ikut dikecualikan dari reset.

    settings.shop_name = current_app.config["CAFE_NAME"]
    settings.address = None
    settings.phone = None
    settings.instagram = None
    settings.tiktok = None
    settings.whatsapp = None
    settings.other_social = None
    settings.wifi_name = None
    settings.wifi_password = None
    settings.app_logo_choice = "square"
    settings.login_logo_choice = "square"
    settings.receipt_logo_choice = "wide"
    settings.navbar_display = "both"
    settings.receipt_paper_width = "58"
    settings.ppn_enabled = False
    settings.ppn_percentage = 11.0
    settings.demo_settings_backup = None
    settings.maintenance_mode = False
    settings.cache_version = (settings.cache_version or 0) + 1

    db.session.commit()


@staff_bp.route("/admin/system/clear-cache", methods=["POST"])
@roles_required(ROLE_OWNER)
def clear_cache():
    settings = get_settings()
    settings.cache_version += 1
    db.session.commit()

    flash(
        _(
            "Cache aplikasi dibersihkan - logo, foto profil, dan QR code "
            "akan otomatis dimuat ulang versi terbaru di semua device."
        ),
        "success",
    )
    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/system/backup/create", methods=["POST"])
@roles_required(ROLE_OWNER)
def backup_create():
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"cafepos_backup_{timestamp}.zip"
    zip_path = os.path.join(_backup_folder(), filename)

    try:
        write_backup_zip(zip_path, _db_file_path(), current_app.static_folder)
    except Exception as exc:  # apa pun penyebabnya, jangan pernah bilang "berhasil"
        current_app.logger.exception("Backup gagal")
        flash(_("Backup GAGAL dibuat: %(error)s", error=str(exc)), "danger")
        return redirect(url_for("staff.admin_settings"))

    flash(_("Backup berhasil dibuat: %(filename)s", filename=filename), "success")
    return redirect(url_for("staff.admin_settings"))


@staff_bp.route("/admin/system/backup/download/<path:filename>")
@roles_required(ROLE_OWNER)
def backup_download(filename):
    filename = secure_filename(filename)
    path = os.path.join(_backup_folder(), filename)

    if not os.path.exists(path):
        abort(404)

    return send_file(path, as_attachment=True, download_name=filename)


@staff_bp.route("/admin/system/backup/delete/<path:filename>", methods=["POST"])
@roles_required(ROLE_OWNER)
def backup_delete(filename):
    filename = secure_filename(filename)
    path = os.path.join(_backup_folder(), filename)

    if os.path.exists(path):
        try:
            os.remove(path)
            flash(_("Backup %(filename)s dihapus.", filename=filename), "success")
        except OSError:
            flash(_("Gagal menghapus file backup (mungkin sedang dipakai)."), "danger")

    return redirect(url_for("staff.admin_settings"))
